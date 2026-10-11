"""Build the Word candidate only from generated matched-experiment CSVs.

Authoring must be invoked after the documents skill's runtime/operation gate.
Render and visually review the candidate before publishing the final DOCX.
"""
import argparse
import csv
import json
from datetime import date
from pathlib import Path

import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parents[1]
BASELINE='FogCloudAStarAdvance';PROPOSED='MistDynamicFogFallback';DYNAMIC='FogCloudDynamicAdvance'
PRIMARY={'route_decision_ms':('Route decision latency','ms'),'ev_response_s':('EV response time','s'),
         'pdr':('Packet delivery ratio','ratio'),'nrl':('Normalized routing load','packets/delivery'),
         'throughput_bps':('EM throughput','bit/s'),'e2e_delay_ms':('EM end-to-end delay','ms'),
         'traffic_light_wait_s':('EV traffic-light waiting','s')}
SECONDARY={'ev_full_stops':('EV full stops','stops/run'),
           'ev_green_approach_fraction':('Green on approach','share'),
           'background_time_loss_s':('Matched background time loss','s/vehicle'),
           'ev_below_1_s':('EV time below 1 m/s','s'),
           'route_changes':('Route changes','events/run'),'fog_requests':('Fog requests','requests/run')}

def rows(path):
    with path.open(newline='',encoding='utf-8') as handle:return list(csv.DictReader(handle))

def table(document,headers,data):
    result=document.add_table(rows=1,cols=len(headers))
    result.style='Light Shading Accent 1'
    for cell,value in zip(result.rows[0].cells,headers):cell.text=value
    marker=OxmlElement('w:tblHeader');result.rows[0]._tr.get_or_add_trPr().append(marker)
    for values in data:
        cells=result.add_row().cells
        for cell,value in zip(cells,values):cell.text=str(value)
    for row in result.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after=Pt(3)
                for run in paragraph.runs:run.font.size=Pt(9)
    document.add_paragraph()
    return result

def chart(summary,metric,label,unit,folder,context):
    densities=list(dict.fromkeys(r['density'] for r in summary))
    configs=[BASELINE,PROPOSED,DYNAMIC]
    if metric=='traffic_light_wait_s':configs.append('NoPreemptionBaseline')
    names={BASELINE:'Baseline: Fog/Cloud advance',PROPOSED:'Proposed: Dynamic Mist + fallback',
           DYNAMIC:'Ablation: Dynamic Fog advance','NoPreemptionBaseline':'No preemption (waiting control)'}
    colors=['#556e8f','#27866b','#c58b2b','#ae5b5b']
    fig,ax=plt.subplots(figsize=(9,4.1))
    width=.78/len(configs)
    for j,config in enumerate(configs):
        group=[next((r for r in summary if r['density']==density and r['configuration']==config and r['metric']==metric),None) for density in densities]
        valid=[(i,r) for i,r in enumerate(group) if r is not None]
        if not valid:continue
        positions=[i-.39+width*(j+.5) for i,r in valid]
        means=[float(r['mean']) for i,r in valid]
        errors=[[max(0,float(r['mean'])-float(r['ci95_lower'])) for i,r in valid],
                [max(0,float(r['ci95_upper'])-float(r['mean'])) for i,r in valid]]
        ax.errorbar(positions,means,yerr=errors,fmt='o',capsize=4,color=colors[j],label=names[config])
    ax.set_xticks(range(len(densities)),densities)
    ax.set_ylabel(f'{label} ({unit})');ax.set_title(f'{context}: {label} [95% CI]')
    ax.grid(axis='y',alpha=.2);ax.legend(fontsize=8,loc='best')
    fig.tight_layout();path=folder/f'{context}-{metric}.png';fig.savefig(path,dpi=150);plt.close(fig)
    return path

def main():
    parser=argparse.ArgumentParser(__doc__)
    parser.add_argument('--output',type=Path,required=True)
    args=parser.parse_args()
    results=ROOT/'results/next_validation'
    audit=json.loads((results/'audit.json').read_text())
    if not audit['passed']:raise SystemExit('Cannot author from unaudited results')
    pairs=rows(results/'paired_comparisons.csv')
    verdicts=rows(results/'requirement_verdicts.csv')
    document=Document()
    section=document.sections[0]
    section.top_margin=section.bottom_margin=Inches(.75)
    section.left_margin=section.right_margin=Inches(.7)
    normal=document.styles['Normal'];normal.font.name='Calibri';normal.font.size=Pt(10)
    normal.paragraph_format.space_after=Pt(6)
    for name,size in (('Title',28),('Heading 1',18),('Heading 2',13)):
        style=document.styles[name];style.font.name='Calibri';style.font.size=Pt(size);style.font.color.rgb=RGBColor.from_string('234F75')
    document.add_paragraph('Emergency Vehicle Navigation',style='Title')
    document.add_paragraph('Matched framework comparison in a VANET',style='Subtitle')
    document.add_paragraph('SUMO 1.18 | OMNeT++ 6.3 | Veins 5.3.1')
    document.add_paragraph('Baseline: FogCloudAStarAdvance\nProposed: MistDynamicFogFallback\nDynamic routing ablation: FogCloudDynamicAdvance')
    document.add_paragraph(date.today().strftime('%d %B %Y'))
    document.add_paragraph('Measured results, project process and explicit limitations')
    document.add_page_break()
    document.add_heading('Supervisor summary',level=1)
    document.add_paragraph('The main comparison gives both frameworks the same advance signal requests. A second Fog control also reroutes dynamically, allowing us to examine whether local Mist computation adds a benefit beyond dynamic routing. Ordinary traffic, heavier demand, fixed blockage and seeded random incidents are reported separately.')
    table(document,['Scenario / density','A','B','C','D'],[[r['scenario']+' / '+r['density'],r['A'],r['B'],r['C'],r['D']] for r in verdicts])
    document.add_paragraph('A needs faster decisions. B uses the previously requested 20% response reduction as a reporting reference. C needs two supported improvements among full stops, green approach share and background time loss. D needs lower, positive signal waiting with a paired interval below zero. No outcome determines which runs are retained.')
    document.add_page_break()
    document.add_heading('How the project works',level=1)
    for heading,text in [
        ('1. Generate shared traffic','SUMO creates the same background trips, incident input and emergency vehicle for every framework within a seed and scenario. Every reported cell contains all seeds 1-30.'),
        ('2. Deliver the emergency message','Roadside units and vehicles exchange the emergency message over the Veins radio model. Delivery failures remain in the results.'),
        ('3. Plan and review the route','Fog/Cloud plans remotely. Mist plans locally and requests Fog after a 500 ms watchdog when initial work has not completed. The dynamic-Fog ablation sends real network requests for periodic route reviews using the same live observations, cost rules and thresholds.'),
        ('4. Request safe signal priority','The matched frameworks use 250 m / 20 s advance requests and identical signal clearance, minimum green and bounded hold rules. The EV still follows SUMO safety and car-following rules.'),
        ('5. Measure and verify','The pipeline reads raw scalars, mobility, signal logs, trip summaries and FCD. Signal states update at 500 ms; metrics accumulate every 100 ms. The output includes confidence intervals, paired sample counts and failed or unfinished cases.')]:
        document.add_heading(heading,level=2);document.add_paragraph(text)
    document.add_paragraph('The 900 ms worker stalls are controlled injected faults. They are not naturally measured hardware failures. No optional queue/signal-aware cost redesign was adopted. Uncongested response gains are constrained by the approximately 123.9 s optimistic external-road lower bound.')
    document.add_page_break()
    document.add_heading('Matched measurements for requirements A-D',level=1)
    document.add_paragraph('Confidence intervals below are paired proposed-minus-baseline bootstrap intervals with 10000 resamples. Negative differences favor proposed for times, stops and background loss; positive differences favor proposed for green share. Percent change is undefined when the baseline is zero. Time comparisons use available matched arrivals; every scheduled run remains in individual CSVs.')
    labels={'route_decision_ms':'A: decision ms','ev_response_s':'B: response s','traffic_light_wait_s':'D: signal wait s',
            'ev_full_stops':'C: full stops','ev_green_approach_fraction':'C: green share','background_time_loss_s':'C: background loss s'}
    for scenario,density in sorted({(r['context'],r['density']) for r in pairs if r['contrast']=='main_matched'}):
        document.add_heading(scenario.replace('_',' ').title()+' / '+density,level=2)
        selected=[r for r in pairs if r['context']==scenario and r['density']==density and r['contrast']=='main_matched' and r['metric'] in labels]
        def f(value):return f'{float(value):.3f}' if value else 'undefined'
        table(document,['Metric','Baseline','Proposed','Change %','95% difference CI','Pairs'],
            [[labels[r['metric']],f(r['baseline_mean']),f(r['proposed_mean']),f(r['percent_change']),
              '['+f(r['ci95_lower'])+', '+f(r['ci95_upper'])+']',r['n_pairs']] for r in selected])
    document.add_heading('Does local computing add a benefit?',level=1)
    document.add_paragraph('Dynamic Fog versus static Fog estimates the value of adding remote dynamic rerouting. Proposed versus dynamic Fog estimates the remaining framework difference. Existing Mist periodic reviews compute synchronously without a second scheduled processing delay, whereas Fog reviews include remote service and radio delay. The contrast also includes initial cloud/Mist processing and controlled fallback; it is not a pure processor benchmark. Lost or stale remote reviews are retained.')
    table(document,['Scenario / density','Contrast','Baseline s','Compared s','Change %','95% difference CI / pairs'],
        [[r['context']+' / '+r['density'],r['contrast'].replace('_',' '),f(r['baseline_mean']),f(r['proposed_mean']),f(r['percent_change']),
          '['+f(r['ci95_lower'])+', '+f(r['ci95_upper'])+'] / '+r['n_pairs']]
         for r in pairs if r['metric']=='ev_response_s' and r['contrast']!='main_matched'])
    document.add_paragraph('Whole-route slow time includes the common scheduled start stop at 30 m. Signal waiting counts speed below 0.1 m/s within 20 m of the next signal. Whole-route time below 1 m/s detects slow motion outside that definition. Zero waiting is kept when measured; no artificial waiting is added. Random incidents can occur after an EV passes the blocked edge.')
    folder=args.output.parent/'matched_word_graphs';folder.mkdir(parents=True,exist_ok=True)
    for context in ('ordinary','dense','fixed_incident','random_incident'):
        path=results/context/'summary.csv'
        if not path.exists():continue
        summary=rows(path)
        for group,metrics in (('Primary metrics',PRIMARY),('Operational metrics',SECONDARY)):
            document.add_page_break();document.add_heading(context.replace('_',' ').title()+': '+group.lower(),level=1)
            for index,(metric,(label,unit)) in enumerate(metrics.items()):
                if index and index%2==0:document.add_page_break()
                image=chart(summary,metric,label,unit,folder,context)
                document.add_picture(str(image),width=Inches(6.8))
                p=document.add_paragraph(label+'; error bars show 95% confidence intervals.')
                p.paragraph_format.space_after=Pt(9)
    document.add_page_break();document.add_heading('Interpretation and reproducibility',level=1)
    document.add_paragraph('The full set of configurations, primary and secondary plots, matched comparisons, pilot results and audit evidence is in results/next_validation and docs/NEXT_RESULTS.md. The predeclared design and any runtime-driven cell reductions are in docs/NEXT_PLAN.md. Historical ordinary and fixed-obstruction evidence is preserved. NoPreemptionBaseline appears only in the signal-waiting comparison.')
    document.add_paragraph('EM throughput represents a fixed emergency-message payload over the measurement window; similar values do not prove that density is unchanged. Mean EV speed depends on journey distance and time and is not an independent performance win. Background time loss uses shared observed vehicle IDs at 900 s; insertion, unfinished trips and cohort counts are disclosed. It is not a pure causal measure of preemption harm.')
    document.add_paragraph('New vector files are losslessly gzip-compressed with uncompressed checksums. The model inputs, seed assignments and binary/source provenance are audited. Reported gains are specific to the measured scenarios; no large gain is promised where there is little congestion or avoidable delay.')
    args.output.parent.mkdir(parents=True,exist_ok=True)
    document.save(args.output)
    print(args.output)

if __name__=='__main__':main()
