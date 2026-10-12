"""Build a unified project report from completed, pipeline-generated evidence.

Render and review the candidate before publishing it as the final submission.
"""
import csv
import json
from pathlib import Path
import sys

import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT/'analysis'))
from next_validation import METRICS

OUT=ROOT/'scratch/complete_project_report.docx'
PILOT=ROOT/'results/next_validation/completed_pilots'
BASE='FogCloudAStarAdvance'
PROP='MistDynamicFogFallback'
FOG='FogCloudDynamicAdvance'
PRIMARY=['pdr','nrl','throughput_bps','e2e_delay_ms','route_decision_ms','ev_response_s','traffic_light_wait_s']
REQUIREMENTS=['route_decision_ms','ev_response_s','ev_full_stops','ev_green_approach_fraction','background_time_loss_s','traffic_light_wait_s']
NAMES={BASE:'Fog advance',PROP:'Dynamic Mist with fallback',FOG:'Dynamic Fog advance',
       'MistAStarAdvance':'Static Mist advance','MistDynamicAStar':'Dynamic Mist',
       'NoPreemptionBaseline':'No preemption'}
EXPLAIN={
 'pdr':'Share of emergency messages delivered. The alert service runs before route calculation, so routing frameworks can have the same delivery ratio.',
 'nrl':'Control and emergency transmissions per delivered alert. It is undefined when no alert is delivered. Lower means less communication overhead.',
 'throughput_bps':'Delivered emergency payload bits divided by the fixed 900 second window. One delivered 256 byte alert contributes 2.27556 bit/s. This is not total network capacity.',
 'e2e_delay_ms':'Time from emergency message generation to EV reception. It measures alert transport before route calculation.',
 'route_decision_ms':'Time from EV alert reception to the first applied route. The pooled result includes the scheduled controlled faults.',
 'ev_response_s':'Time from alert generation until the EV reaches the accident. Driving, queueing and signals dominate this measure.',
 'traffic_light_wait_s':'EV time below 0.1 m/s within 20 m of the next signal. A measured zero is retained. This definition does not include all crawling or remote queues.',
 'ev_full_stops':'Stops below 0.1 m/s after reaching 1 m/s, with recovery to 1 m/s between stops. The common scheduled start stop is included.',
 'ev_below_1_s':'Whole-route time below 1 m/s after navigation release. This reveals crawling that the narrow signal-wait definition may miss.',
 'ev_whole_stopped_s':'Whole-route standstill time, including stops outside the 20 m signal zone and the shared scheduled start stop.',
 'ev_green_approach_fraction':'Share of first observed approaches within 20 m with a green link lamp. A red arrival followed by green crossing remains a red approach.',
 'ev_green_fraction':'Share of observed signal entries with the prior approach lamp green. Safe red-light compliance can make this high for all frameworks.',
 'ev_green_count':'Number of observed green signal entries. Route choice can change the number of signals encountered.',
 'background_time_loss_s':'Mean SUMO time loss for common observed background vehicle IDs at 900 seconds, including unfinished observed trips. It is not a pure causal estimate of preemption harm.',
 'ev_mean_speed_mps':'Distance divided by travel time. This is useful context but depends on journey time and is not counted as an independent win.',
 'route_changes':'Applied replacement routes. A route review without a replacement does not count as a change.',
 'route_reviews':'Logged route evaluations. Reviews can skip an internal junction or an invalid remaining route, so this is not a count of every timer firing.',
 'fallback_triggered':'Share of scheduled runs that actually requested Fog after the watchdog. The 900 ms stalls are controlled injected faults, not organic failures.',
 'fallback_decision_ms':'Time to apply the fallback route in runs where fallback actually occurred. Missing values do not mean zero latency.',
 'control_transmissions':'Recorded control packet transmissions per run. These are supporting network-load measurements.',
 'control_bytes':'Recorded control traffic bytes per run. The graph retains the original metric definition.',
 'ev_distance_m':'Measured EV route distance. A longer diversion may still be faster when it avoids a queue.',
 'ev_travel_s':'Time from EV departure to accident arrival. This overlaps response time and is not an independent performance win.',
 'ev_delay_vs_freeflow_s':'Travel time minus measured distance divided by 13.9 m/s, bounded at zero. This is correlated with response time.',
 'fog_requests':'Actual Fog route requests, including remote reviews in the dynamic Fog ablation. Lower counts can reflect the framework design.',
 'fog_completed_reviews':'Remote dynamic route reviews with an accepted reply. Lost or stale reviews are retained in the evidence.',
 'fog_unanswered_reviews':'Remote review requests without a completed accepted reply. These disclose communication and review limitations.'}

def read(path):
    with path.open(newline='',encoding='utf-8') as stream:return list(csv.DictReader(stream))

def num(value):return f'{float(value):.3f}' if value not in ('',None) else 'undefined'

def metric_name(key):return METRICS[key][0]

def heading(doc,text):doc.add_heading(text,level=1)

def para(doc,text):doc.add_paragraph(text)

def table(doc,headers,values,widths=None):
    t=doc.add_table(rows=1,cols=len(headers));t.autofit=False
    widths=widths or [6.7/len(headers)]*len(headers)
    for col,width in zip(t.columns,widths):col.width=Inches(width)
    for cell,value in zip(t.rows[0].cells,headers):cell.text=str(value)
    marker=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(marker)
    for row in values:
        for cell,value in zip(t.add_row().cells,row):cell.text=str(value)
    for i,row in enumerate(t.rows):
        for j,cell in enumerate(row.cells):
            cell.width=Inches(widths[j]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            pr=cell._tc.get_or_add_tcPr()
            shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'203E56' if i==0 else ('EFF3F6' if i%2 else 'FFFFFF'));pr.append(shade)
            borders=OxmlElement('w:tcBorders')
            for side in ('top','left','bottom','right'):
                line=OxmlElement('w:'+side);line.set(qn('w:val'),'single');line.set(qn('w:sz'),'4');line.set(qn('w:color'),'D9D9D9');borders.append(line)
            pr.append(borders)
            for p in cell.paragraphs:
                p.paragraph_format.space_after=Pt(4);p.paragraph_format.space_before=Pt(4)
                if j:p.alignment=WD_ALIGN_PARAGRAPH.CENTER
                for run in p.runs:run.font.size=Pt(9);run.bold=i==0;run.font.color.rgb=RGBColor.from_string('FFFFFF' if i==0 else '000000')
    para(doc,'')

def paired_table(doc,rows,pilot=False):
    values=[]
    labels={'route_decision_ms':'A Decision ms','ev_response_s':'B Response s',
            'ev_full_stops':'C Full stops','ev_green_approach_fraction':'C Green approach share',
            'background_time_loss_s':'C Background loss s','traffic_light_wait_s':'D Signal wait s'}
    for r in rows:
        baseline=r['baseline_mean'] if pilot else r['mean_B']
        proposed=r['proposed_mean'] if pilot else r['mean_A']
        values.append([r['density'].replace('heavy','')+' '+labels[r['metric']],num(baseline),num(proposed),
                       num(r['percent_change']) if r['percent_change'] else 'n/a','['+num(r['ci95_lower'])+', '+num(r['ci95_upper'])+']',r['n_pairs']])
    table(doc,['Density and metric','Baseline','Proposed','Change %','Paired 95% interval','Pairs'],values,[1.8,.9,.9,.8,1.8,.5])

def graph(summary,key,context,folder,pilot=False):
    densities=[d for d in ('low','medium','high','heavy400','heavy600') if any(r['density']==d for r in summary)]
    configs=[BASE,'MistAStarAdvance','MistDynamicAStar',PROP] if not pilot else [BASE,PROP,FOG]
    if key=='traffic_light_wait_s':configs+=['NoPreemptionBaseline']
    colors={BASE:'#556e8f','MistAStarAdvance':'#c58b2b','MistDynamicAStar':'#8d6bad',
            PROP:'#27866b',FOG:'#8b5e3c','NoPreemptionBaseline':'#ae5b5b'}
    fig,ax=plt.subplots(figsize=(9,3.8))
    for j,config in enumerate(configs):
        points=[]
        for i,density in enumerate(densities):
            r=next((r for r in summary if r['density']==density and r['configuration']==config and r['metric']==key),None)
            if r is not None and r['mean'] not in ('',None):points.append((i,r))
        if not points:continue
        positions=[i-.34+.68*(j+.5)/len(configs) for i,r in points]
        means=[float(r['mean']) for i,r in points]
        errors=[[max(0,float(r['mean'])-float(r['ci95_lower'])) for i,r in points],
                [max(0,float(r['ci95_upper'])-float(r['mean'])) for i,r in points]]
        ax.errorbar(positions,means,yerr=errors,fmt='o',capsize=3,color=colors[config],label=NAMES[config])
    label,unit=METRICS[key]
    ax.set_xticks(range(len(densities)),[d.replace('heavy','') for d in densities]);ax.set_xlabel('Traffic demand')
    ax.set_ylabel(label+' ('+unit+')');ax.set_title(context+' '+label+' [95% CI]')
    ax.grid(axis='y',alpha=.2);ax.legend(fontsize=7,loc='best');fig.tight_layout()
    path=folder/(context.replace(' ','_')+'-'+key+'.png');fig.savefig(path,dpi=150);plt.close(fig)
    return path

def gallery(doc,summary,context,folder,pilot=False):
    available={r['metric'] for r in summary}
    order=PRIMARY+[k for k in METRICS if k not in PRIMARY]
    keys=[k for k in order if k in available and any(r['metric']==k and r['configuration'] in (BASE,PROP,FOG,'MistAStarAdvance','MistDynamicAStar') for r in summary)]
    for i,key in enumerate(keys):
        if i%2==0:doc.add_page_break();heading(doc,context+' graphs')
        doc.add_picture(str(graph(summary,key,context,folder,pilot)),width=Inches(6.7))
        p=doc.add_paragraph(EXPLAIN[key]+' Error bars show 95 percent confidence intervals.');p.paragraph_format.space_after=Pt(8)
    return len(keys)

def main():
    audit=json.loads((PILOT/'audit.json').read_text());assert audit['passed'] and audit['complete_runs']==109
    doc=Document();s=doc.sections[0];s.page_width=Inches(8.27);s.page_height=Inches(11.69)
    s.top_margin=s.bottom_margin=Inches(.65);s.left_margin=s.right_margin=Inches(.78)
    for style in doc.styles:
        if style.type in (1,2):
            style.font.name='Liberation Sans'
            pr=style.element.find(qn('w:pPr'))
            if pr is not None:
                for border in list(pr.findall(qn('w:pBdr'))):pr.remove(border)
    normal=doc.styles['Normal'];normal.font.name='Liberation Sans';normal.font.size=Pt(10)
    normal.paragraph_format.space_after=Pt(7)
    for name,size in [('Title',28),('Subtitle',15),('Heading 1',18),('Heading 2',13)]:
        style=doc.styles[name];style.font.name='Liberation Sans';style.font.size=Pt(size);style.font.color.rgb=RGBColor(0,0,0)
    doc.add_paragraph('Emergency Vehicle Navigation in a Vehicular Network',style='Title')
    doc.add_paragraph('Complete project report and measured results',style='Subtitle')
    para(doc,'SUMO 1.18   OMNeT++ 6.3   Veins 5.3.1')
    para(doc,'Baseline FogCloudAStarAdvance\nProposed MistDynamicFogFallback')
    para(doc,'11 October 2026')
    para(doc,'Project overview, system design, experimental method, matched results, completed pilot evidence and limitations')
    doc.add_page_break();heading(doc,'Project overview')
    para(doc,'An emergency vehicle needs to learn where an accident happened and reach it quickly while other traffic continues to move safely. This project simulates emergency message delivery, road route planning and traffic signal priority on a small road grid. It compares remote Fog and Cloud computation with local Mist computation, dynamic route reviews and Fog recovery when the initial local work stalls.')
    para(doc,'The strongest complete comparison gives baseline and proposed the same advance signal request policy. Proposed makes the first route decision faster. Its ordinary-traffic response improvement is only 0.3 to 0.7 percent. In the fixed road obstruction it is about 40 to 45 percent faster, with fewer full stops and lower matched background time loss. A consistently lower and positive signal waiting time is not established.')
    para(doc,'The later heavy-demand experiment was stopped before new reported seeds began. This report includes its 109 fully completed held-out pilot runs and excludes 11 unfinished or unstarted cases. The completed pilot response means do not favor proposed. Their limited coverage cannot establish a final congestion or local-computing claim.')
    heading(doc,'Guide to the report')
    for text in ['System design and the complete operating process','Routing frameworks and fair comparisons','Experimental settings and metric definitions','Complete ordinary and fixed obstruction results','Original signal policy and dynamic Fog ablations','Completed heavy-demand pilots and coverage','Conclusions and requirements that remain unmet','Graph gallery for every available matched and pilot metric']:
        doc.add_paragraph(text,style='List Bullet')
    doc.add_page_break();heading(doc,'System design')
    para(doc,'The road network has 16 signalized junctions in a 4 by 4 grid and 48 directed road links. Three roadside units exchange radio messages with vehicles within the configured 400 m neighborhood. SUMO moves vehicles and operates the signals. OMNeT++ schedules the communication and application events. Veins connects the radio simulation to SUMO through TraCI.')
    table(doc,['Component','Role'],[['SUMO','Roads, vehicles, car following, traffic signals and trip records'],['OMNeT++ and Veins','Wireless events and connection to road movement'],['Roadside units','Relay emergency alerts and road observations'],['Mist routing','Local route calculation from received observations'],['Fog and Cloud','Remote route computation and fallback service'],['Metrics collector','Timing, delivery, mobility, signal and routing evidence']],[1.6,5.1])
    heading(doc,'How an emergency is handled')
    for text in ['Generate shared background trips for the chosen density and seed, and insert the emergency and accident vehicles. Every framework receives the same requested traffic inputs within that scenario and seed.',
                 'Create a 256 byte emergency alert at the accident. Vehicles and roadside units relay it through the radio network. A failed delivery remains a failed case in the results.',
                 'Hold the EV until its first route is ready. Fog routing requests a remote result. Mist calculates locally. Proposed requests Fog if initial Mist work remains unfinished after 500 ms.',
                 'Apply the route and release the EV under normal SUMO safety rules. The EV respects car following, acceleration, junction priority and red lights, with a 13.9 m/s speed cap.',
                 'Send advance priority requests within 250 m or an estimated 20 seconds of the next signal in the matched frameworks. The signal controller preserves minimum green, yellow and all-red clearance.',
                 'Dynamic frameworks review received traffic information while driving. Replace the route only when its cost and stability rules allow a sufficiently better alternative.',
                 'Record arrival, signal approaches, whole-route slow motion, route reviews, route changes and communication counts. Process the raw records through the existing analysis pipeline and retain failures.']:
        doc.add_paragraph(text,style='List Number')
    doc.add_page_break();heading(doc,'Routing frameworks')
    table(doc,['Configuration','Computation and review','Signal requests'],[
        [BASE,'Remote initial A star route without periodic rerouting','Advance 250 m or 20 s'],
        ['MistAStarAdvance','Local initial A star route without periodic rerouting','Advance 250 m or 20 s'],
        ['MistDynamicAStar','Local A star with periodic live route reviews','Advance 250 m or 20 s'],
        [PROP,'Dynamic local routing and watchdog Fog recovery','Advance 250 m or 20 s'],
        [FOG,'Remote routing and remote dynamic route reviews','Advance 250 m or 20 s'],
        ['FogCloudAStar and MistAStar','Retained original static frameworks','Reactive 100 m or 8 s'],
        ['NoPreemptionBaseline','Fog routing with priority disabled','Waiting control only']],[1.85,3.15,1.7])
    para(doc,'A star searches for a low-cost route using a free-flow heuristic. Dynamic routing uses received beacons and roadside edge reports to estimate current travel cost; it does not know future traffic. At least three observed vehicles are needed for a congestion adjustment. Speeds are capped at the road speed limit in the cost estimate.')
    para(doc,'Frozen settings review every five seconds. A replacement needs 10 percent lower remaining cost, or a confirmed slow edge with five percent improvement. Two slow observations and a 30 second gap reduce oscillation. Reviews and applied replacements are separate measurements. Held-out seeds 101 to 110 were used for tuning; reported seeds 1 to 30 were not used for selection. The reference settings were retained. An optional signal-aware cost redesign was not adopted.')
    para(doc,'Initial Mist processing is modeled as 300 ms plus 2 ms per expanded A star node. Periodic local reviews are synchronous simulation calculations. Dynamic Fog reviews include remote service and radio latency. This is a framework comparison and not a hardware benchmark of an actual onboard computer.')
    doc.add_page_break();heading(doc,'Signal priority and fallback')
    para(doc,'A compatible green can be extended. Otherwise the controller waits for the active phase to receive its ten second minimum green, then uses two seconds of yellow and one second of all-red clearance. Signal processing takes 150 ms. Priority green lasts at least ten seconds. The hold ends after EV passage or at the 25 second maximum. Retries are allowed every five seconds, but duplicate requests do not extend an active hold.')
    para(doc,'Advance requests can prepare green before the EV arrives. A genuine zero signal wait is therefore possible and is not replaced with an invented minimum. The approach trace records signal distance, link lamp and controller phase. Whole-route time below 1 m/s and full stops expose slow motion outside the 20 m waiting zone. The shared vehicle input also has a scheduled stop at 30 m; whole-route stop measures include it.')
    para(doc,'The initial watchdog is 500 ms and telemetry validation adds 0 ms. A fixed independently sampled set of reported seeds receives a controlled 900 ms local worker stall in every Mist framework: 4, 5, 9, 10, 13, 17, 18 and 29. These are injected faults, not organic fallback or measured resource contention. Normal and controlled-stall cohorts remain separate in the CSV exports. A selected seed with no delivered alert cannot activate fallback.')
    para(doc,'Separate short-notice red-signal validation uses 18 completed stress runs. It verifies positive measured waiting and protected transitions with deliberately late priority requests. These tests validate the mechanism and are not pooled with the ordinary or fixed-obstruction results.')
    doc.add_page_break();heading(doc,'Experimental design')
    table(doc,['Setting','Value'],[['Road and radio','4 by 4 grid, 48 directed links, three RSUs, 400 m neighborhood'],['Ordinary demand','72 low, 144 medium and 200 high background trips over 360 s'],['Reported seeds','All scheduled seeds 1 to 30 for each complete cell'],['Heavy-demand pilots','400 and 600 trips; 360 s and 120 s departure windows; seeds 101 to 110'],['Observation horizon','900 s per completed pilot run'],['EV maximum speed','13.9 m/s, with SUMO safety mode 31'],['Timing','500 ms watchdog, 0 ms telemetry delay, 100 ms polling'],['Underlying road states','500 ms SUMO movement updates']],[1.65,5.05])
    para(doc,'The complete matched-policy extension contains 630 ordinary runs and 540 fixed-obstruction runs across its configuration cells. The differing counts reflect the scheduled waiting-control coverage. Earlier original-policy experiments and isolated signal tests remain preserved separately. No unfinished simulation contributes a metric to the completed pilot export.')
    para(doc,'The fixed obstruction requests three stopped vehicles on C1C2 at 90, 91 and 92 seconds, with stops until 250 seconds. SUMO safety can delay insertion. FCD verifies a physical blocker during the incident. All frameworks share the same requested incident inputs. This is one fixed road scenario, not evidence for all incident locations.')
    para(doc,'Demand seeds vary trip generation and OMNeT++ random streams. SUMO driving randomness is held common at seed zero by the Veins launch behavior. Different demand seeds still change routes and entry edges. Scheduled trip totals differ from vehicles simultaneously on the road. FCD and trip summaries verify actual queues and observed traffic.')
    para(doc,'No baseline was slowed or weakened. Input changes apply to every framework within a scenario. Every scheduled reported seed remains in the individual CSVs, including nondeliveries. New random-incident reported results were not produced before cancellation and are not presented as measured evidence.')
    doc.add_page_break();heading(doc,'Seven primary metrics')
    table(doc,['Metric and unit','Meaning'],[[metric_name(k)+' ('+METRICS[k][1]+')',EXPLAIN[k]] for k in PRIMARY],[1.85,4.85])
    para(doc,'Route decision latency, EM delay and fallback latency are displayed in milliseconds. Raw precision is retained; displayed tables round values to three decimals. EV response and signal waiting remain in seconds. Percent change is 100 times proposed minus baseline divided by baseline. Negative values favor proposed for time, delay, load and stop counts; positive green share is favorable. A zero baseline gives an undefined percentage, shown as n/a.')
    heading(doc,'Intervals and missing observations')
    para(doc,'The complete policy experiment retains its existing methods: Wilson intervals for proportions, percentile bootstrap intervals for signal waiting, and Student t intervals for other continuous metrics. Its paired comparisons use paired Student t intervals. Completed pilots retain their pipeline bootstrap means and paired differences with 10000 resamples and fixed seed 42, and Wilson proportion intervals. Graph titles use the generic 95 percent CI label. No method was changed to create a favorable result.')
    para(doc,'PDR includes scheduled runs. Delivery-dependent measures use available deliveries or matched arrivals, with sample counts shown. NRL is undefined without a delivered alert. The second pilot has incomplete cell coverage, so its completed subset is descriptive and may be biased by completion patterns. Intervals crossing zero do not establish a directional improvement. Results are exploratory and not adjusted for multiple comparisons.')
    summaries={};policy_pairs={}
    for context in ('ordinary','obstruction'):
        summaries[context]=read(ROOT/f'results/policy_validation/{context}/summary.csv')
        policy_pairs[context]=read(ROOT/f'results/policy_validation/{context}/paired_comparisons.csv')
        doc.add_page_break();heading(doc,('Ordinary traffic' if context=='ordinary' else 'Fixed obstruction')+' primary results')
        para(doc,'The main comparison holds advance signal requests equal. Each entry below is mean [95 percent interval], followed by sample count. These are complete reported experiments, not the later pilots.')
        for density in ('low','medium','high'):
            doc.add_heading(density.title()+' demand',level=2)
            vals=[]
            for k in PRIMARY:
                cells=[metric_name(k)+' ('+METRICS[k][1]+')']
                for config in (BASE,'MistAStarAdvance','MistDynamicAStar',PROP):
                    r=next(r for r in summaries[context] if r['configuration']==config and r['density']==density and r['metric']==k)
                    cells.append(num(r['mean'])+'\n['+num(r['ci95_lower'])+', '+num(r['ci95_upper'])+']\nn='+r['n'])
                vals.append(cells)
            table(doc,['Metric','Fog advance','Static Mist advance','Dynamic Mist','Proposed'],vals,[1.55,1.29,1.29,1.29,1.28])
        doc.add_page_break();heading(doc,('Ordinary traffic' if context=='ordinary' else 'Fixed obstruction')+' requirements A to D')
        for density in ('low','medium','high'):
            rows=[r for r in policy_pairs[context] if r['comparison']=='matched_policy' and r['density']==density and r['metric'] in REQUIREMENTS]
            paired_table(doc,rows)
        para(doc,'A is route decision latency. B is EV response time. C considers full stops, green approach share and matched background time loss. D requires lower and positive signal waiting. Intervals are proposed minus baseline in the metric unit; sample counts are matched pairs.')
        para(doc,'Ordinary traffic supports faster decisions but only a minor response reduction, no two additional supported operational wins, and zero signal waiting in both main frameworks.' if context=='ordinary' else 'Fixed obstruction supports a clear response reduction and two additional improvements in full stops and background time loss. Waiting is zero for proposed at low demand, and the paired waiting intervals cross zero at every density. Requirement D remains unmet.')
    doc.add_page_break();heading(doc,'Separating signal policy and routing effects')
    para(doc,'The original FogCloudAStar baseline requests priority at 100 m or eight seconds. Proposed requests at 250 m or 20 seconds. That original comparison mixes computation, routing and signal timing. FogCloudAStarAdvance is the fair headline baseline because it shares the advance policy. The following preserved response comparisons show the difference.')
    vals=[]
    for context in ('ordinary','obstruction'):
        for density in ('low','medium','high'):
            for comp in ('original_policy','matched_policy','fog_signal_policy'):
                r=next((r for r in policy_pairs[context] if r['density']==density and r['comparison']==comp and r['metric']=='ev_response_s'),None)
                if r:vals.append([context+' '+density,comp.replace('_',' '),num(r['mean_B']),num(r['mean_A']),num(r['percent_change']),r['n_pairs']])
    table(doc,['Scenario','Comparison','Reference s','Compared s','Change %','Pairs'],vals,[1.25,1.5,1.05,1.05,1.1,.75])
    para(doc,'Small or equal PDR and EM throughput do not demonstrate broken routing. These metrics describe the same alert transport before routing. Live rerouting matters when there is avoidable road delay. On the small ordinary grid the trip is close to its physical floor: external roads alone need at least 1721.6/13.9 = 123.856 seconds, before junctions, acceleration and computation. Against the matched baseline, a 20 percent ordinary improvement is not physically supported by this road bound.')
    doc.add_page_break();heading(doc,'Completed heavy demand pilots')
    para(doc,'The first pilot used 400 and 600 trips over a 360 second departure window. All 60 runs completed. Baseline response rose from 148.25 to 150.00 seconds, only 1.75 seconds, missing the predeclared five second congestion criterion despite observed queues. The second pilot concentrated departures into 120 seconds, with shared inputs across frameworks. It was cancelled after 49 completed runs. No new reported seeds 1 to 30 were started.')
    coverage=read(PILOT/'coverage.csv')
    table(doc,['Window s','Trips','Framework','Complete','Scheduled'],[[r['pilot_window_s'],r['density'].replace('heavy',''),NAMES[r['configuration']],r['completed_runs'],r['scheduled_runs']] for r in coverage],[.75,.65,3.4,.95,.95])
    para(doc,'Included seed lists and excluded cases are disclosed in coverage.csv. Every included run has scalar, compressed vector, index and normal 900 second termination evidence; vector checksums are verified. Eleven unfinished or unstarted cases contribute no values. No missing dynamic Fog 600-trip value is invented for the second pilot.')
    pairs=read(PILOT/'paired_comparisons.csv')
    for window in (360,120):
        doc.add_page_break();heading(doc,f'Completed pilot results with {window} second departures')
        rows=[r for r in pairs if r['contrast']=='main_matched' and int(r['pilot_window_s'])==window and r['metric'] in REQUIREMENTS]
        paired_table(doc,rows,pilot=True)
        para(doc,'Completed pilot response means are higher for proposed at both trip levels. A response win is not established. Signal waiting and secondary measures are mixed; use the paired intervals and arrival counts. These held-out observations are not a substitute for a completed reported experiment.')
    doc.add_page_break();heading(doc,'Does local computing add a benefit')
    para(doc,'FogCloudDynamicAdvance adds remote route reviews using the same dynamic rules and advance requests. Comparing it with static Fog estimates the effect of remote rerouting. Comparing proposed with dynamic Fog examines the remaining local framework difference, including initial processing and controlled fallback. Synchronous local reviews and delayed remote requests prevent interpretation as a pure processor benchmark.')
    rows=[r for r in pairs if r['metric']=='ev_response_s' and r['contrast']!='main_matched']
    table(doc,['Window and trips','Comparison','Reference s','Compared s','Paired 95% interval','Pairs'],[[r['pilot_window_s']+' / '+r['density'].replace('heavy',''),r['contrast'].replace('_',' '),num(r['baseline_mean']),num(r['proposed_mean']),'['+num(r['ci95_lower'])+', '+num(r['ci95_upper'])+']',r['n_pairs']] for r in rows],[1.0,1.7,.85,.85,1.7,.6])
    para(doc,'The complete first pilot and completed subset of the second pilot are retained even where proposed is slower. The absent second-pilot 600-trip dynamic Fog cell leaves the stronger-demand ablation unfinished. The available evidence cannot establish a general local-computing advantage beyond dynamic rerouting.')
    doc.add_page_break();heading(doc,'Conclusions and remaining requirements')
    table(doc,['Requirement','Ordinary matched','Fixed obstruction matched','Heavy-demand pilots'],[
        ['A Faster route decision','Supported at every density','Supported at every density','Lower completed means; held-out evidence only'],
        ['B Clear response margin','Only 0.3 to 0.7 percent; 20 percent reference unmet','About 40 to 45 percent faster','Not established; proposed response means higher'],
        ['C Two additional improvements','Not established','Full stops and background time loss improve','Mixed results; no final reported verdict'],
        ['D Lower nonzero signal waiting','Both zero; unmet','Paired intervals cross zero; unmet','Mixed and incomplete; unmet']],[1.55,1.7,1.7,1.75])
    para(doc,'The project demonstrates an operating message-to-route-to-signal workflow and a useful dynamic-routing response under the specified fixed obstruction. It does not meet every supervisor requirement. The ordinary route has little avoidable delay, the fixed incident is only one scenario, and the newer demand study is incomplete. No measurement was changed to create a large visual difference or a positive waiting time.')
    heading(doc,'Reproducing and reading the evidence')
    para(doc,'The source, analysis and tests are in the EmergencyNavigation repository. Complete matched CSVs are under results/policy_validation/ordinary and obstruction. Completed pilot CSVs, coverage and validity audit are under results/next_validation/completed_pilots. POLICY_VALIDATION.md documents the complete matched results, COMPLETED_PILOT_RESULTS.md the completed pilots, and METHODOLOGY.md the calculation methods. The earlier submission is preserved separately as Emergency_Vehicle_Navigation_Historical_e3178f2.docx.')
    para(doc,'The graphs below are generated from the pipeline summary CSVs. Graph means use each configuration\'s available observations. Paired tables use only common valid seeds, so their means and counts can differ when arrivals are missing. Every available metric is included, even if proposed loses. Primary metrics come first, followed by secondary data. Error bars always retain 95 percent intervals; a narrow interval may be visually small. NoPreemptionBaseline is shown only for signal waiting. Ordinary, fixed obstruction and the two pilot windows are separate sections. Raw logs, CSVs and graphs are not edited by hand.')
    folder=ROOT/'scratch/complete_project_graphs';folder.mkdir(exist_ok=True)
    graph_count=0
    for context in ('ordinary','obstruction'):
        graph_count+=gallery(doc,summaries[context],'Ordinary traffic' if context=='ordinary' else 'Fixed obstruction',folder)
    for window in (360,120):
        graph_count+=gallery(doc,read(PILOT/f'pilot_{window}/summary.csv'),f'Pilot {window} s',folder,True)
    OUT.parent.mkdir(exist_ok=True);doc.save(OUT)
    (ROOT/'scratch/complete_project_report_manifest.json').write_text(json.dumps({'graphs':graph_count,'pilot_runs':109,'excluded_cases':11,'tables':len(doc.tables),'output':str(OUT)},indent=2))
    print(json.dumps({'output':str(OUT),'graphs':graph_count,'tables':len(doc.tables)}))

if __name__=='__main__':main()
