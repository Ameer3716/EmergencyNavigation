#!/usr/bin/env python3
"""Generate a standalone, plain-English project report from measured CSVs."""
from __future__ import annotations

import csv
import shutil
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT.parent / "EmergencyNavigation_Reviewed_400m_Submission" / "Emergency_Vehicle_Navigation_Final_Submission.docx"
CONFIGS = ("FogCloudAStar", "MistAStar", "MistDynamicAStar", "MistDynamicFogFallback")
SHORT = ("Fog/Cloud", "Mist", "Mist Dynamic", "Mist + Fog")
DENSITIES = ("low", "medium", "high")
VEHICLES = {"low": 72, "medium": 144, "high": 200}
PRIMARY = (
    ("pdr", "PDR", "ratio"),
    ("nrl", "NRL", "packets/delivery"),
    ("throughput_bps", "EM throughput", "bit/s"),
    ("e2e_delay_ms", "EM end-to-end delay", "ms"),
    ("route_decision_ms", "Route decision latency", "ms"),
    ("ev_response_s", "EV response time", "s"),
    ("traffic_light_wait_s", "EV traffic-light waiting time", "s"),
)
GRAPH_INFO = (
    ("pdr", "Packet delivery ratio", "The share of generated emergency messages that reached the EV. All four approaches share this alert service before route computation; equal bars can be correct."),
    ("nrl", "Normalized routing load", "Control and EM transmissions per delivered emergency message. Lower means less communication overhead."),
    ("throughput_bps", "EM throughput", "Delivered EM payload bits divided by the fixed 900 second observation window. One delivery contributes 2.27556 bit/s. This measures the shared alert service."),
    ("e2e_delay_ms", "EM end-to-end delay", "Time from RSU message generation to EV reception, in milliseconds. It measures the shared alert service before any route computation."),
    ("route_decision_ms", "Route decision latency", "Time from EV message reception to the initial route being applied, in milliseconds."),
    ("ev_response_s", "EV response time", "Time from emergency message generation until the EV reaches the accident, in seconds."),
    ("traffic_light_wait_s", "EV traffic-light waiting time", "EV standstill near a signal stop line, in seconds. A value of zero means no measured wait."),
    ("fallback_triggered", "Fallback activation rate", "Share of Mist with Fog runs where the 500 ms watchdog requested Fog. Controlled stalls are explicitly injected in eight of the 30 seeds."),
    ("route_changes", "Applied route changes", "Mean number of actual congestion or cost reroutes per run. Reviews without a replacement are excluded."),
    ("route_reviews", "Route computation frequency", "Mean number of periodic Dynamic A* route evaluations per run, normally scheduled every five seconds."),
    ("fallback_decision_ms", "Fallback decision latency", "Time to apply the Fog route in runs where the watchdog actually triggered, in milliseconds."),
    ("control_transmissions", "Control transmissions", "Mean number of control-packet transmissions per run. This is supporting communication-load evidence."),
    ("control_bytes", "Control bytes", "Mean control traffic volume in bytes per run. This is supporting communication-load evidence."),
    ("ev_delay_vs_freeflow_s", "EV corridor delay versus free flow", "Extra travel time relative to the modeled free-flow corridor, in seconds."),
    ("ev_distance_m", "EV route distance", "Distance traveled by the EV to reach the accident, in metres."),
    ("ev_travel_s", "EV travel time", "EV movement time from departure to accident arrival, in seconds."),
)
PROCESS_STEPS = (
    "Create a small road network with 16 junctions and traffic lights. Put three roadside units (RSUs) beside the roads. Each RSU can communicate within a 400 m radio neighborhood.",
    "Add normal traffic to the roads. The low, medium, and high traffic scenarios contain 72, 144, and 200 background vehicles. Thirty random seeds give different trips for each scenario.",
    "Start the emergency event. An accident vehicle stops, and the nearby RSU sends one 256-byte emergency message through the vehicle and roadside network to the emergency vehicle (EV).",
    "When the EV receives the message, calculate a route to the accident using the selected routing configuration. Some configurations use a fixed A* route; the dynamic configurations review live road costs every five seconds.",
    "While the EV drives under SUMO safety rules, Fog/Cloud and static Mist request signal priority within 100 m or eight seconds. Dynamic Mist requests earlier, within 250 m or 20 seconds, to prepare the junction before arrival. Requests may retry every five seconds. Active holds reject duplicates and remain bounded at 25 seconds. The controller preserves a ten-second minimum green, then yellow and all-red clearance.",
    "Mist normally computes the route with no added telemetry validation delay. Eight independently selected seeds receive a controlled 900 ms Mist stall in all three Mist approaches. In Mist with Fog, unfinished work after 500 ms triggers a Fog request. These are labeled fault tests, not naturally occurring failures.",
    "Record whether the message arrived, how long communication and route decisions took, how the EV traveled, and how many reviews, route changes, and fallback events occurred. Each run has a 900 second observation window.",
    "Repeat the four primary configurations at three traffic levels and 30 matched seeds: 360 runs. Add 90 no-preemption control runs, displayed only for traffic-light waiting. Calculate averages and 95% confidence intervals, report normal and controlled-stall results separately, and draw the graphs.",
)


def read(path: Path) -> list[dict[str, str]]:
    with path.open(newline="", encoding="utf-8") as handle:
        return list(csv.DictReader(handle))


def fmt(value: float, key: str) -> str:
    if key == "pdr":
        return f"{value:.3f}"
    if key == "nrl":
        return f"{value:,.0f}"
    if key == "throughput_bps":
        return f"{value:.3f}"
    if key in ("ev_response_s", "traffic_light_wait_s"):
        return f"{value:.3f}"
    return f"{value:.1f}"


def graph_value(value: float, key: str) -> str:
    if key in ("pdr", "fallback_triggered", "throughput_bps"):
        return f"{value:.3f}"
    if key in ("route_changes", "route_reviews"):
        return f"{value:.2f}"
    if key in ("nrl", "control_transmissions", "control_bytes"):
        return f"{value:,.0f}"
    if key in ("ev_response_s", "ev_travel_s", "traffic_light_wait_s", "ev_delay_vs_freeflow_s"):
        return f"{value:.3f}"
    return f"{value:.1f}"


def graph_blocks(summary: dict[tuple[str, str, str], dict[str, str]]) -> list[tuple[str, str, str, Path, str]]:
    blocks = []
    for key, label, explanation in GRAPH_INFO:
        for density in DENSITIES:
            path = ROOT / "results/graphs" / f"{key}-{density}.png"
            if not path.exists():
                raise SystemExit(f"Missing graph: {path}")
            values = []
            labels = list(zip(CONFIGS, SHORT))
            if key == "traffic_light_wait_s":
                labels.append(("NoPreemptionBaseline", "No preemption"))
            for cfg, short in labels:
                row = summary.get((cfg, density, key))
                if row is not None:
                    values.append(f"{short} {graph_value(float(row['mean']), key)}")
            caption = (f"{density.capitalize()} traffic ({VEHICLES[density]} background vehicles): "
                       + "; ".join(values) + ". Bars show means and 95% confidence intervals.")
            blocks.append((key, label, explanation, path, caption))
    cohorts = read(ROOT / 'results/processed/summary-cohorts.csv')
    pairs = read(ROOT / 'results/processed/paired_comparisons.csv')
    for condition, condition_label in (("normal", "Normal operation"), ("controlled_stall", "Controlled Mist stalls")):
        for metric, label, unit in (("route_decision_ms", "Route decision latency", "ms"), ("ev_response_s", "EV response time", "s")):
            for density in DENSITIES:
                rows = [r for r in cohorts if r['scenario_condition'] == condition and r['density'] == density and r['metric'] == metric]
                caption = f"{density.capitalize()} traffic, {condition_label.lower()}: " + "; ".join(
                    f"{SHORT[CONFIGS.index(r['configuration'])]} {float(r['mean']):.2f} {unit} (n={r['n']})" for r in rows) + ". Error bars show 95% confidence intervals."
                blocks.append((f"{metric}_{condition}", f"{condition_label} {label.lower()}",
                               "These figures separate seeds without a stall from seeds assigned the controlled 900 ms Mist stall. Fog/Cloud is unaffected by that local stall.",
                               ROOT / f"results/graphs/{metric}-{condition}-{density}.png", caption))
    for metric, label, unit in (("ev_response_s", "EV response time", "s"), ("route_decision_ms", "Route decision latency", "ms")):
        for density in DENSITIES:
            rows = [r for r in pairs if r['density'] == density and r['metric'] == metric]
            caption = f"{density.capitalize()} traffic: " + "; ".join(
                f"{SHORT[CONFIGS.index(r['config_A'])]} minus {SHORT[CONFIGS.index(r['config_B'])]} {float(r['mean_paired_diff']):.2f} {unit}, 95% CI [{float(r['ci95_lower']):.2f}, {float(r['ci95_upper']):.2f}], n={r['n_pairs']}" for r in rows) + "."
            blocks.append((f"paired_{metric}", f"Matched seed differences in {label.lower()}",
                           "Each bar subtracts configuration B from A for the same delivered seeds. Below zero means A is faster. An interval crossing zero does not establish a difference; these comparisons are exploratory and not adjusted for multiple testing.",
                           ROOT / f"results/graphs/paired_{metric}-{density}.png", caption))
    if {b[3].name for b in blocks} != {p.name for p in (ROOT / "results/graphs").glob("*.png")}:
        raise SystemExit("The report gallery must include every current graph exactly once")
    supplemental = ROOT / 'results/supplemental'
    if (supplemental / 'validation_report.json').exists():
        rows = read(supplemental / 'route_transaction_summary.csv') + read(supplemental / 'red_signal_summary.csv')
        extras = (
            ('route_requests_per_run', 'Fog route requests', 'Mean wireless Fog route requests per scheduled run. Local Mist needs no Fog request. These counts are separate from the emergency alert.', 'requests/run'),
            ('route_transaction_ms', 'Fog route transaction turnaround', 'Time from sending a Fog request to accepting its route reply. Processing and Cloud backhaul are included; watchdog waiting is excluded. Local-only Mist has no wireless transaction.', 'ms'),
            ('route_network_roundtrip_ms', 'Fog route network round trip', 'Request and reply network time after subtracting processing and Cloud backhaul. The 500 ms watchdog wait is also excluded.', 'ms'),
            ('red_signal_wait_s', 'Short notice red signal validation', 'Isolated red-signal scenarios use two matched seeds per density. Priority is deliberately requested late, within 10 metres or half a second. These stress tests are separate from the main comparison.', 's'),
        )
        for metric, label, explanation, unit in extras:
            for density in DENSITIES:
                values = [r for r in rows if r['metric'] == metric and r['density'] == density and r['mean']]
                caption = f'{density.capitalize()} traffic: ' + '; '.join(
                    f"{('No preemption' if r['configuration'] == 'NoPreemptionBaseline' else SHORT[CONFIGS.index(r['configuration'])])} {float(r['mean']):.3f} {unit} (n={r['n']})" for r in values) + '. Error bars show 95% confidence intervals.'
                blocks.append((metric, label, explanation, supplemental / f'graphs/{metric}-{density}.png', caption))
    incident = ROOT / 'results/congestion_validation'
    if (incident / 'validation_report.json').exists():
        rows = read(incident / 'summary.csv') + read(incident / 'fault_recovery_summary.csv')
        for metric, label, explanation, unit in (
            ('ev_response_s', 'Controlled incident response', 'A separate matched experiment adds the same stopped queue after the initial route decision. These values measure response under that incident, not ordinary traffic.', 's'),
            ('route_changes', 'Controlled incident route changes', 'Actual route replacements based on received traffic reports in the incident experiment. No routing costs or route choices are hardcoded.', 'changes/run'),
            ('recovery_saved_ms', 'Controlled fault route decision time saved', 'Matched time saved by Fog fallback against Dynamic Mist under the same controlled 900 ms stall in the original batch. This is route readiness, not journey time.', 'ms')):
            for density in DENSITIES:
                values = [r for r in rows if r['metric'] == metric and r['density'] == density]
                caption = f'{density.capitalize()} traffic: ' + '; '.join(
                    f"{SHORT[CONFIGS.index(r['configuration'])]} {float(r['mean']):.3f} {unit} (n={r['n']})" for r in values) + '. Error bars show 95% confidence intervals.'
                blocks.append((f'incident_{metric}', label, explanation, incident / f'graphs/{metric}-{density}.png', caption))
        pairs = read(incident / 'paired_comparisons.csv')
        for density in DENSITIES:
            values = [r for r in pairs if r['metric'] == 'ev_response_s' and r['density'] == density]
            caption = f'{density.capitalize()} traffic: ' + '; '.join(
                f"{SHORT[CONFIGS.index(r['config_A'])]} minus {SHORT[CONFIGS.index(r['config_B'])]} {float(r['mean_paired_diff']):.3f} s, 95% CI [{float(r['ci95_lower']):.3f}, {float(r['ci95_upper']):.3f}], n={r['n_pairs']}" for r in values) + '.'
            blocks.append(('incident_paired_response', 'Controlled incident matched response differences',
                           'Negative means the first algorithm is faster. The paired interval quantifies uncertainty. A controlled obstacle does not establish the same benefit in ordinary traffic.',
                           incident / f'graphs/incident_paired_response_s-{density}.png', caption))
    review = ROOT / 'results/client_review'
    if (review / 'provenance.json').exists():
        pairs = read(review / 'baseline_comparisons.csv')
        for context in ('Ordinary traffic', 'Controlled obstruction'):
            for metric, label, explanation in (
                ('route_decision_ms', 'Direct baseline route decision latency', 'Time from alert reception to the initial route, in milliseconds. Lower is better.'),
                ('ev_response_s', 'Direct baseline EV response time', 'Time from alert generation to arrival, in seconds. Lower is better. The framework comparison includes both routing and signal-request policy.'),
                ('traffic_light_wait_s', 'Direct baseline signal waiting', 'Standstill near a signal stop line, in seconds. Lower is better; a genuine green arrival may have zero waiting.'),
                ('ev_delay_vs_freeflow_s', 'Excess travel delay above free flow', 'Actual EV travel time minus traveled distance divided by 13.9 m/s, bounded at zero. Lower is better. This supporting metric is correlated with response time.'),
                ('fog_route_requests', 'Fog route requests per run', 'Actual initial RouteRequest send events at the EV. Lower means fewer remote route transactions. This is not a count of relays or EM delivery.')):
                group = [r for r in pairs if r['context'] == context and r['metric'] == metric and r['config_A'] == 'MistDynamicFogFallback']
                caption = context + ': ' + '; '.join(
                    f"{r['density'].capitalize()} Fog {float(r['mean_B']):.3f}, proposed {float(r['mean_A']):.3f}, reduction {float(r['reduction_percent']):.2f}%, paired n={r['n_pairs']}"
                    if r['reduction_percent'] else f"{r['density'].capitalize()}: baseline zero; percentage undefined"
                    for r in group) + '. Error bars show 95% confidence intervals.'
                path = review / 'graphs' / f'{metric}-{context.lower().replace(" ", "_")}.png'
                blocks.append((f'baseline_{metric}_{context}', label, explanation, path, caption))
    return blocks


def write_graph_guide(blocks: list[tuple[str, str, str, Path, str]]) -> None:
    lines = ["# Emergency vehicle navigation project and graph guide", "",
             f"This project studies how an emergency vehicle receives an accident alert and chooses a route through traffic. The Word report contains the complete project explanation and the same {len(blocks)} graphs.", "",
             "## How the system works", ""]
    for index, step in enumerate(PROCESS_STEPS, 1):
        lines += [f"{index}. {step}", ""]
    lines += ["## How to read the figures", "",
              "Each bar is a density and configuration mean. Error bars show 95% confidence intervals; the exact methods are in [METHODOLOGY.md](METHODOLOGY.md). Metrics requiring EM delivery use only valid delivered runs. PDR, EM throughput, and fallback activation use all 30 scheduled runs where applicable. Fallback latency uses only runs where the watchdog actually activated.", "",
              f"## All {len(blocks)} graphs", ""]
    previous = None
    for key, label, explanation, path, caption in blocks:
        if key != previous:
            lines += [f"### {label}", "", explanation, ""]
            previous = key
        lines += [f"![{label} for {path.stem.rsplit('-', 1)[1]} traffic](../{path.relative_to(ROOT).as_posix()})", "",
                  caption, ""]
    lines += ["The complete numeric means, valid sample sizes, standard deviations, and intervals are in `../results/processed/summary-batch.csv`. The seven main measures are also tabulated in `VERIFICATION.md`.", ""]
    (ROOT / "docs/PROCESS_AND_GRAPHS.md").write_text("\n".join(lines), encoding="utf-8")


def shade(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def table(doc: Document, headers: tuple[str, ...], body: list[tuple[str, ...]]) -> None:
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = True
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        border = OxmlElement(f"w:{edge}")
        border.set(qn("w:val"), "single")
        border.set(qn("w:sz"), "4")
        border.set(qn("w:color"), "D9D9D9")
        borders.append(border)
    t._tbl.tblPr.append(borders)
    tr_pr = t.rows[0]._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    tr_pr.append(repeat)
    for i, title in enumerate(headers):
        cell = t.rows[0].cells[i]
        cell.text = title
        shade(cell, "1F3957")
        for run in cell.paragraphs[0].runs:
            run.font.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)
            run.font.size = Pt(8)
    for index, values in enumerate(body):
        row = t.add_row()
        for i, value in enumerate(values):
            cell = row.cells[i]
            cell.text = str(value)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if index % 2:
                shade(cell, "F2F6FA")
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                for run in paragraph.runs:
                    run.font.size = Pt(8)
    doc.add_paragraph()


def heading(doc: Document, title: str, level: int = 1) -> None:
    doc.add_heading(title.replace('-', ' '), level=level)


def main() -> None:
    summary_rows = read(ROOT / "results/processed/summary-batch.csv")
    individual = read(ROOT / "results/processed/individual_runs-batch.csv")
    fallback = read(ROOT / "results/processed/fallback_validation.csv")
    paired = read(ROOT / "results/processed/paired_comparisons.csv")
    if len(individual) != 450:
        raise SystemExit("The 360 primary runs and 90 waiting controls are required")
    summary = {(r["configuration"], r["density"], r["metric"]): r for r in summary_rows}
    cohort = {(r['scenario_condition'], r['configuration'], r['density'], r['metric']): r
              for r in read(ROOT / 'results/processed/summary-cohorts.csv')}
    blocks = graph_blocks(summary)
    real_fallback = [r for r in fallback if r["configuration"] == "MistDynamicFogFallback" and r["fallback_triggered"] == "1"]
    if not real_fallback:
        raise SystemExit("No real batch fallback found")

    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin = section.bottom_margin = Inches(0.7)
    section.left_margin = section.right_margin = Inches(0.65)
    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(9.5)
    normal.font.color.rgb = RGBColor(25, 25, 25)
    normal.paragraph_format.space_after = Pt(6)
    for style_name in ("Title", "Heading 1", "Heading 2"):
        style = doc.styles[style_name]
        style.font.name = "Arial"
        style.font.color.rgb = RGBColor(0, 0, 0)
    title_properties = doc.styles["Title"].element.get_or_add_pPr()
    for border in title_properties.findall(qn("w:pBdr")):
        title_properties.remove(border)
    doc.styles["Title"].font.size = Pt(20)
    doc.styles["Heading 1"].font.size = Pt(14)
    doc.styles["Heading 2"].font.size = Pt(11)

    cover_title = doc.add_paragraph(style="Title")
    cover_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cover_title.paragraph_format.space_before = Pt(105)
    cover_title.add_run("Emergency Vehicle Navigation in a Vehicular Network")
    cover_subtitle = doc.add_paragraph()
    cover_subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cover_subtitle.paragraph_format.space_before = Pt(18)
    cover_subtitle.add_run("A simulation study of Fog and Mist routing with dynamic updates and fallback").italic = True
    cover_details = doc.add_paragraph()
    cover_details.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cover_details.paragraph_format.space_before = Pt(105)
    cover_details.add_run("SUMO + OMNeT++ + Veins\n")
    cover_details.add_run("4 routing configurations  |  3 traffic levels  |  30 seeds\n")
    cover_details.add_run(f"360 primary runs + 90 waiting controls  |  {len(blocks)} result graphs\n")
    if (ROOT / 'results/supplemental/validation_report.json').exists():
        cover_details.add_run('18 separate red signal validation runs\n')
    if (ROOT / 'results/congestion_validation/validation_report.json').exists():
        cover_details.add_run('360 separate controlled incident runs\n')
    cover_details.add_run("October 2026")
    doc.add_page_break()

    heading(doc, "Project overview")
    doc.add_paragraph(
        "An emergency vehicle needs to reach an accident quickly. It must first learn where the "
        "accident is, then choose a road route while normal traffic keeps moving. This project "
        "tests four ways to make that route decision in a simulated vehicular network."
    )
    doc.add_paragraph(
        "The project measures alert delivery, routing work, and the emergency vehicle's journey. "
        "The routing comparison tests the location of route computation, live traffic updates, "
        "and Fog recovery. All four approaches share the alert delivery mechanism and traffic-light "
        "priority. The alert arrives before route computation, so its delivery ratio, throughput "
        "and delay can legitimately be equal across routing approaches."
    )

    direct = read(ROOT / 'results/client_review/baseline_comparisons.csv')
    proposed_comparisons = [r for r in direct if r['config_A'] == 'MistDynamicFogFallback']
    heading(doc, 'Main measured findings')
    doc.add_paragraph('The clearest journey advantage appears when a queue develops after the first route decision. '
                      'This is the situation in which live routing has useful new information. '
                      'The separate ordinary traffic experiment shows the size of the benefit when that obstruction is absent. '
                      'The following results compare the complete proposed Mist with Fog recovery framework against Fog and Cloud.')
    findings = []
    for context, metric, label in (
        ('Controlled obstruction', 'ev_response_s', 'Response time with obstruction'),
        ('Ordinary traffic', 'ev_response_s', 'Response time in ordinary traffic'),
        ('Ordinary traffic', 'route_decision_ms', 'Route decision latency'),
        ('Ordinary traffic', 'ev_delay_vs_freeflow_s', 'Excess travel delay'),
        ('Ordinary traffic', 'fog_route_requests', 'Fog route requests'),
    ):
        group = [r for r in proposed_comparisons if r['context'] == context and r['metric'] == metric]
        reductions = [float(r['reduction_percent']) for r in group]
        findings.append((label, f'{min(reductions):.1f} to {max(reductions):.1f}% lower'))
    table(doc, ('Measure', 'Reduction across the three densities'), findings)
    doc.add_paragraph('Ordinary signal waiting falls from baseline means of less than two seconds to zero. '
                      'Under obstruction, mean waiting also falls at every density, but the paired confidence intervals '
                      'establish a clear reduction only at high density. Zero waiting can occur when advance requests '
                      'prepare a green before arrival; separate late request tests verify that the safety controller can still require waiting. '
                      'Excess travel delay is related to response time, while Fog request counts measure a separate communication requirement.')
    doc.add_paragraph('A 20 percent response improvement is demonstrated in the obstruction experiment. '
                      'The ordinary experiment does not meet that target. The target was chosen after the results were measured, '
                      'so it is used to describe the effect size rather than as a predeclared statistical test. '
                      'All scheduled seeds and both experiments remain in the evidence.')

    heading(doc, "The simulated road and communication system")
    doc.add_paragraph(
        "The road is a 4 by 4 grid of signalized junctions. SUMO moves the vehicles and operates "
        "the traffic lights. OMNeT++ and Veins simulate wireless messages between vehicles and "
        "three roadside units (RSUs). The RSUs have a 400 m communication neighborhood. The "
        "accident vehicle, RSUs, and emergency vehicle exchange the information needed to "
        "start the response."
    )
    doc.add_paragraph(
        "Normal traffic comes in three levels: 72 generated trips in low traffic, 144 in medium "
        "traffic, and 200 in high traffic, spread over 360 seconds. These are total demand counts; "
        "fewer vehicles may be on the road at any one time. Thirty demand seeds create different "
        "routes and entry edges. The same demand and OMNeT++ seed are used for each routing approach. "
        "Veins launchd holds SUMO driving randomness at seed zero."
    )

    heading(doc, "How the system works step by step")
    for step in PROCESS_STEPS:
        doc.add_paragraph(step, style="List Number")

    heading(doc, "The four routing approaches")
    table(doc, ("Approach", "Plain-English description"), [
        ("FogCloudAStar", "Fog and cloud services calculate the first A* route. The EV follows that route without periodic route reviews."),
        ("MistAStar", "The EV's nearby Mist layer calculates the first A* route. There are no periodic route reviews."),
        ("MistDynamicAStar", "Mist calculates the route and checks live road costs every five seconds. A better, stable route can replace the current route."),
        ("MistDynamicFogFallback", "Dynamic Mist routing runs normally. If the initial Mist decision exceeds the 500 ms watchdog, Fog calculates the route instead."),
    ])
    doc.add_paragraph(
        "A* is a shortest-path search. The dynamic versions use recent vehicle and roadside "
        "messages to estimate changing road costs. A route review is only a check; it counts "
        "as a route change when the EV actually receives a replacement route."
    )
    doc.add_paragraph('The dynamic frameworks combine those route reviews with advance signal requests at 250 metres or 20 seconds. Fog/Cloud and static Mist use reactive requests at 100 metres or eight seconds. Every preempting approach can retry a junction after five seconds. The same controller preserves minimum green and clearance, rejects duplicates during an active hold, and releases priority after the EV passes or the 25-second limit. Earlier preparation can avoid braking as well as standstill waiting. The comparison measures these complete frameworks, so a journey improvement cannot be attributed to A* computation alone.')

    heading(doc, "What was measured")
    table(doc, ("Main measure", "What it tells us"), [
        ("Packet delivery ratio (PDR)", "The share of accident alerts that reached the emergency vehicle."),
        ("Normalized routing load (NRL)", "The number of message transmissions per delivered accident alert."),
        ("EM throughput", "Delivered emergency-message payload bits per second of the fixed observation window."),
        ("EM end-to-end delay", "Milliseconds from sending the accident alert to its arrival at the emergency vehicle."),
        ("Route decision latency", "Milliseconds from alert reception until the first route is applied."),
        ("EV response time", "Seconds from alert generation until the emergency vehicle reaches the accident."),
        ("EV traffic-light waiting time", "Seconds the emergency vehicle waits near a signal stop line."),
    ])
    doc.add_paragraph(
        "We also counted Fog fallback activations, actual route changes, five-second route "
        "reviews, and the time taken to apply a Fog route after fallback. Supporting measures also include "
        "excess travel delay above free flow and Fog route requests per run. Each graph shows an "
        "average with a 95% confidence interval."
    )

    heading(doc, "Main results")
    primary_runs = [r for r in individual if r['configuration'] in CONFIGS]
    delivered = sum(int(r['delivered_messages']) for r in primary_runs)
    doc.add_paragraph(
        f"The accident alert reached the emergency vehicle in {delivered} of the 360 primary runs. "
        "The delivery and timing outcomes reflect this grid, its radio connections, and its "
        "traffic trips. A higher traffic count does not guarantee a particular direction of change in every measure."
    )
    doc.add_paragraph(
        "The tables give the average, its 95% confidence interval in brackets, and the "
        "number of valid runs. Delivery ratio and EM throughput use all 30 runs at each "
        "traffic level. Travel and timing measures use runs in which the alert arrived."
    )
    for density in DENSITIES:
        heading(doc, f"{density.capitalize()} traffic", 2)
        data = []
        for key, label, unit in PRIMARY:
            cells = []
            for cfg in CONFIGS:
                r = summary[(cfg, density, key)]
                cells.append(f"{fmt(float(r['mean']), key)}\n[{fmt(float(r['ci95_lower']), key)}, {fmt(float(r['ci95_upper']), key)}]\nn={r['n']}")
            data.append((f"{label}\n({unit})", *cells))
        table(doc, ("Metric", *SHORT), data)

    heading(doc, "What the main results mean")
    doc.add_paragraph(
        "The alert is delivered before the EV calculates its route. PDR and EM end-to-end "
        "delay therefore measure the same radio event in every routing approach. EM "
        "throughput is calculated from those same deliveries and a fixed payload, so it "
        "also can match across approaches. These values should change with density when "
        "the number of successful deliveries changes; they do not have to change just "
        "because the route calculator is different."
    )
    heading(doc, "How the measured results change with traffic density", 2)
    density_rows = []
    for d in DENSITIES:
        dynamic_runs = [r for r in individual if r['density'] == d and r['configuration'] == 'MistDynamicAStar']
        peak = sum(float(r['peak_active_background']) for r in dynamic_runs) / len(dynamic_runs)
        pdr = float(summary['MistAStar', d, 'pdr']['mean'])
        throughput = float(summary['MistAStar', d, 'throughput_bps']['mean'])
        nrl = sum(float(summary[c, d, 'nrl']['mean']) for c in CONFIGS) / len(CONFIGS)
        response = float(summary['MistDynamicAStar', d, 'ev_response_s']['mean'])
        density_rows.append((d.capitalize(), str(VEHICLES[d]), f"{peak:.1f}", f"{pdr:.3f}",
                             f"{throughput:.3f}", f"{nrl:,.0f}", f"{response:.2f}"))
    table(doc, ("Traffic", "Trips", "Mean peak vehicles", "PDR", "EM throughput (bit/s)",
                "Mean NRL", "Dynamic Mist response (s)"), density_rows)
    doc.add_paragraph(
        "This shows real density effects: generated demand rises from 72 to 200 trips, "
        "and the table reports observed PDR, throughput and peak active background vehicles. Throughput is "
        "similar across routing approaches within each density because their alert delivery "
        "counts match; it is not constant across all three densities."
    )
    heading(doc, "What changes between routing approaches", 2)
    response_rows = []
    for d in DENSITIES:
        changes = sum(int(float(r['route_changes'])) for r in individual
                      if r['density'] == d and r['configuration'] == 'MistDynamicAStar')
        response_rows.append((d.capitalize(),
            f"{float(summary['MistAStar', d, 'ev_response_s']['mean']):.2f}",
            f"{float(summary['MistDynamicAStar', d, 'ev_response_s']['mean']):.2f}",
            f"{float(summary['MistDynamicFogFallback', d, 'ev_response_s']['mean']):.2f}",
            str(changes)))
    table(doc, ("Traffic", "Mist A* response (s)", "Dynamic Mist response (s)",
                "Dynamic Mist + Fog response (s)", "Dynamic Mist route changes"), response_rows)
    doc.add_paragraph(
        "Mist A* and Dynamic Mist can have the same initial decision latency because both "
        "use the same Mist A* calculation. Dynamic Mist then reviews live road costs every "
        "five seconds; the logs show applied route changes in the table. The response-time "
        "means and matched intervals quantify whether a difference is supported. "
        "Dynamic frameworks also use advance signal requests, so journey differences do not isolate routing alone."
    )
    for density in DENSITIES:
        response = [float(summary[c,density,'ev_response_s']['mean']) for c in CONFIGS]
        doc.add_paragraph(f"In {density} traffic, the four average EV response times range from "
                          f"{min(response):.2f} to {max(response):.2f} seconds. Read the intervals "
                          "alongside these averages; a small difference alone does not prove a general advantage.")
        for a, b in (("MistDynamicAStar", "MistAStar"), ("MistDynamicFogFallback", "MistDynamicAStar")):
            p = next(r for r in paired if r['density'] == density and r['config_A'] == a and r['config_B'] == b and r['metric'] == 'ev_response_s')
            other = {r['seed']: r['ev_response_s'] for r in individual if r['configuration'] == b and r['density'] == density and r['ev_response_s']}
            identical = sum(r['ev_response_s'] == other.get(r['seed']) for r in individual
                            if r['configuration'] == a and r['density'] == density and r['ev_response_s'])
            doc.add_paragraph(f"For matched {density} seeds, {SHORT[CONFIGS.index(a)]} minus {SHORT[CONFIGS.index(b)]} "
                              f"averages {float(p['mean_paired_diff']):.2f} s, with a 95% interval of "
                              f"{float(p['ci95_lower']):.2f} to {float(p['ci95_upper']):.2f} s. "
                              f"{identical} of {p['n_pairs']} matched response times are identical. "
                              "The paired difference graphs show the uncertainty in these exploratory comparisons.")

    heading(doc, "Normal operation and controlled stalls")
    doc.add_paragraph(
        "The main averages include 22 normal seeds and eight independently selected fault-test seeds. "
        "To show how the routing approaches behave in each condition, the next tables separate "
        "their initial decision times. All Mist approaches face the same injected stall in the "
        "selected seeds. The experiment does not estimate how often such a fault occurs on real hardware."
    )
    initial_nodes, all_nodes = [], []
    for run in individual:
        if not run['configuration'].startswith('Mist'):
            continue
        path = ROOT / f"artifacts/logs/batch/routing-{run['configuration']}-{run['density']}-seed{run['seed']}.csv"
        for event in read(path) if path.exists() else []:
            if event['action'] in ('computed', 'evaluated'):
                all_nodes.append(int(event['expandedNodes']))
            if event['action'] == 'computed' and event['reason'] == 'initial':
                initial_nodes.append(int(event['expandedNodes']))
    initial_range = str(min(initial_nodes)) if min(initial_nodes) == max(initial_nodes) else f"{min(initial_nodes)} to {max(initial_nodes)}"
    doc.add_paragraph(f"The initial calculations expanded {initial_range} A* nodes. "
                      f"The maximum across initial calculations and later reviews was {max(all_nodes)} nodes. "
                      f"The model assigns 300 ms plus 2 ms per node, giving {300+2*max(all_nodes)} ms at that maximum, "
                      f"which leaves {500-(300+2*max(all_nodes))} ms before the watchdog. "
                      "Matching expansion counts can give matching normal decision times across traffic levels. "
                      "Pooled means also depend on which normal and controlled-stall runs received the alert.")
    for density in DENSITIES:
        heading(doc, density.capitalize() + " traffic decision time", 2)
        body = []
        for condition, label in (("normal", "Normal seeds"), ("controlled_stall", "Controlled stall seeds")):
            body.append((label, *(f"{float(cohort[condition,c,density,'route_decision_ms']['mean']):.2f} ms"
                                  for c in CONFIGS)))
        table(doc, ("Condition", *SHORT), body)

    heading(doc, "Traffic-light waiting with and without priority")
    doc.add_paragraph(
        "The no-preemption control uses Fog/Cloud routing with priority disabled. It appears "
        "only in waiting comparisons and their graphs. The priority controller "
        "rechecks the active phase, preserves its minimum green, and uses clearance before serving the EV. Some "
        "requests can finish before arrival; others require the EV to wait. Waiting is measured "
        "from movement rather than added artificially."
    )
    table(doc, ("Traffic", *SHORT, "No preemption"), [
        (d.capitalize(), *(f"{float(summary[c,d,'traffic_light_wait_s']['mean']):.2f} s"
                           for c in CONFIGS + ("NoPreemptionBaseline",))) for d in DENSITIES])

    heading(doc, "Fallback and route changes")
    body = []
    for density in DENSITIES:
        group = [r for r in individual if r["density"] == density and r["configuration"] == "MistDynamicFogFallback"]
        count = sum(int(float(r["fallback_triggered"])) for r in group)
        dynamic = [r for r in individual if r["density"] == density and r["configuration"] == "MistDynamicAStar"]
        body.append((density.capitalize(), f"{count}/30 ({count/30:.1%})",
                     str(sum(int(float(r["route_changes"])) for r in dynamic)),
                     str(sum(int(float(r["route_reviews"])) for r in dynamic)),
                     str(sum(int(float(r["route_changes"])) for r in group)),
                     str(sum(int(float(r["route_reviews"])) for r in group))))
    table(doc, ("Traffic", "Fog takeovers", "Dynamic Mist route changes", "Dynamic Mist reviews", "Mist + Fog route changes", "Mist + Fog reviews"), body)
    example = real_fallback[0]
    activations = {d: sum(int(float(r['fallback_triggered'])) for r in individual
                         if r['configuration'] == 'MistDynamicFogFallback' and r['density'] == d) for d in DENSITIES}
    doc.add_paragraph(
        f"Fog took over in {sum(activations.values())} of 90 Mist with Fog runs: "
        f"{activations['low']} in low, {activations['medium']} in medium, and {activations['high']} "
        "in high traffic. The controlled 900 ms stall delays initial Mist work beyond the "
        "500 ms watchdog. For example, in "
        f"{example['density']} traffic seed {example['seed']}, Fog applied the route "
        f"after {float(example['final_decision_latency_ms']):.2f} ms."
    )
    doc.add_paragraph(f"That example separates {float(example['wait_before_fog_s'])*1000:.2f} ms of watchdog waiting, "
                      f"{float(example['fog_computation_s'])*1000:.2f} ms of Fog computation, and "
                      f"{float(example['communication_delay_s'])*1000:.3f} ms of network communication. "
                      "Watchdog waiting is not counted as communication delay.")
    doc.add_paragraph(
        "The table counts actual route changes separately from five-second reviews. A "
        "review that finds no sufficiently better route keeps the existing route. The "
        "routing logs record both the current route cost and the candidate cost, so these "
        "decisions can be checked from the evidence."
    )

    heading(doc, "Why EM throughput is similar across routing approaches")
    throughput_means = ', '.join(f"{float(summary['MistAStar', d, 'throughput_bps']['mean']):.3f}" for d in DENSITIES)
    doc.add_paragraph(
        "The accident alert has a fixed 256-byte useful payload. Throughput divides the "
        "delivered alert bits by the same 900 second window in every run. One successful "
        "delivery therefore contributes 2.27556 bit/s, regardless of how many normal "
        "vehicles are on the road. This measure describes emergency-message delivery over "
        "the observation window; it is not the total capacity of the wireless network."
    )
    doc.add_paragraph(
        "Traffic does change with density: the scenarios contain 72, 144, or 200 normal "
        "vehicles; trip starting points differ between random seeds; and the dynamic route "
        "logs contain different live road costs and route changes. The throughput means "
        f"for static Mist are {throughput_means} bit/s from low to high density. Within each density, "
        "the four routing approaches have similar throughput because they delivered the "
        "same number of alerts."
    )

    supplemental = ROOT / 'results/supplemental'
    if (supplemental / 'validation_report.json').exists():
        heading(doc, 'Supplementary communication and red signal checks')
        doc.add_paragraph('The seven headline measures keep their original definitions. The additional route transaction measures count Fog requests and accepted replies after the emergency alert. Local Mist uses no wireless Fog transaction, so its wireless transaction delay and completion rate are not applicable. Equal alert delivery results can coexist with different route communication needs.')
        doc.add_paragraph('Route transaction turnaround includes processing and Cloud backhaul. The network round trip subtracts those components and excludes watchdog waiting. These values come from the original raw batch, without inventing unrecorded route-packet byte counts.')
        rows = read(supplemental / 'route_transaction_summary.csv')
        lookup = {(r['configuration'], r['density'], r['metric']): r for r in rows}
        body = []
        for density in DENSITIES:
            for cfg in ('FogCloudAStar', 'MistDynamicFogFallback'):
                r = lookup[cfg, density, 'route_completion_rate']
                body.append((density.capitalize(), SHORT[CONFIGS.index(cfg)], r['requests_sent'], r['replies_accepted'], f"{float(r['mean'])*100:.1f}%"))
        table(doc, ('Traffic', 'Approach', 'Requests', 'Accepted replies', 'Completion'), body)
        doc.add_paragraph('Completion is conditional on sending a Fog request; it is not EM PDR and does not mean every scheduled run received the emergency alert.')
        doc.add_paragraph('Eighteen separate signal tests reuse the existing binary and matched trips. Junction A1 starts with a protected conflicting phase for 150 seconds. Its permissive green is removed in the validation input, keeping both ambulance movements from A0A1 red. Priority requests are deliberately late, within 10 metres or half a second. Fog/Cloud, Mist with Fog, and no-preemption are compared at all three densities with seeds 1 and 4 over 400 seconds. These controlled cases test a short-notice boundary condition; they are not the default priority algorithm and are not added to the main comparison.')
        signals = read(supplemental / 'red_signal_summary.csv')
        table(doc, ('Traffic', 'Approach', 'Runs', 'Mean waiting s'), [(r['density'].capitalize(), 'No preemption' if r['configuration'] == 'NoPreemptionBaseline' else SHORT[CONFIGS.index(r['configuration'])], r['n'], f"{float(r['mean']):.3f}") for r in signals])
        doc.add_paragraph('Every tested priority case has an observed red approach, positive measured waiting, confirmed arrival, and less waiting than its matched no-preemption case. Yellow, all-red, and minimum-green timing are checked from the event logs. Two seeds per density are mechanism validation; wide confidence intervals are retained. Ordinary green-arrival cases may still have zero waiting.')

    incident = ROOT / 'results/congestion_validation'
    if (incident / 'validation_report.json').exists():
        heading(doc, 'Routing under a developing queue')
        doc.add_paragraph('A separate 360-run comparison tests what happens when the chosen road becomes blocked after the first route decision. The input requests three passenger vehicles on C1C2 at 90, 91 and 92 seconds, with stops until 250 seconds. Every algorithm uses the same incident input, background trips, 30 seeds per density, binary, safety controller and controlled stall assignment. Dynamic Mist uses advance priority requests while Fog/static Mist uses reactive requests. SUMO safety checks may delay insertion. These three requested vehicles are additional to ordinary demand. Received beacons and RSU reports provide congestion information; the algorithm receives no privileged incident notification.')
        doc.add_paragraph('Routing retains its minimum of three observed vehicles for a congestion adjustment; incident and ordinary queued vehicles can supply those observations. FCD verifies at least one stopped blocker through the end of the planned incident in every case. Initial route decisions precede the incident, and every delivered-alert run reaches the destination. A pilot uses seeds 1, 4, 22 and 27 at each density to check the mechanism before the full matrix. This controlled test does not describe how often incidents occur or guarantee improvements in normal traffic.')
        rows = read(incident / 'summary.csv')
        table(doc, ('Traffic', 'Approach', 'Arrived samples', 'Mean response s'),
              [(r['density'].capitalize(), SHORT[CONFIGS.index(r['configuration'])], r['n'], f"{float(r['mean']):.3f}")
               for r in sorted((r for r in rows if r['metric'] == 'ev_response_s'),
                               key=lambda r: (DENSITIES.index(r['density']), CONFIGS.index(r['configuration'])))])
        incident_runs = read(incident / 'individual_runs.csv')
        full_queue = sum(int(r['maximum_stopped_incident_vehicles']) == 3 for r in incident_runs)
        exposure_matches = sum(len({(r['maximum_stopped_incident_vehicles'], r['incident_vehicles_inserted_after_stop_period'])
                                    for r in incident_runs if r['density'] == d and int(r['seed']) == s}) == 1
                               for d in DENSITIES for s in range(1, 31))
        doc.add_paragraph(f'{full_queue} runs have all three incident vehicles stopped during the planned period. '
                          f'Realised obstruction and late insertion counts match across the four approaches in {exposure_matches} of 90 density and seed groups. '
                          'All 360 cases are retained with the same requested inputs. Signal and routing policies can change subsequent traffic and insertion states; the comparison includes those consequences. No scenario input was changed after measuring this variation. Every run still has a physical blocker through the end of the incident. Per-vehicle observations and per-configuration counts are provided in REQUIREMENTS_EVIDENCE.md and validation_report.json.')
        avoided = []
        for density in DENSITIES:
            for cfg in CONFIGS[2:]:
                group = [r for r in incident_runs if r['density'] == density and r['configuration'] == cfg]
                avoided.append((density.capitalize(), SHORT[CONFIGS.index(cfg)],
                                sum(int(r['incident_avoided_by_reroute']) for r in group), len(group)))
        doc.add_paragraph('The following counts use actual applied routes that remove the obstructed edge during the incident. A review without a replacement does not count.')
        table(doc, ('Traffic', 'Dynamic approach', 'Runs avoiding the queue', 'Scheduled runs'), avoided)
        pairs = read(incident / 'paired_comparisons.csv')
        table(doc, ('Traffic', 'Matched comparison', 'Difference s', '95% interval s'),
              [(r['density'].capitalize(), f"{SHORT[CONFIGS.index(r['config_A'])]} minus {SHORT[CONFIGS.index(r['config_B'])]}",
                f"{float(r['mean_paired_diff']):.3f}", f"{float(r['ci95_lower']):.3f} to {float(r['ci95_upper']):.3f}")
               for r in pairs if r['metric'] == 'ev_response_s'])
        heading(doc, 'What fallback changes under a controlled fault')
        recovered = read(incident / 'fault_recovery_summary.csv')
        table(doc, ('Traffic', 'Matched stalled cases', 'Route decision time saved ms'),
              [(r['density'].capitalize(), r['n'], f"{float(r['mean']):.3f}") for r in recovered])
        doc.add_paragraph('The fallback savings compare actual initial route application times in the original batch under the same 900 ms Mist stall. This isolates recovery performance from ordinary traffic and does not imply the same saving in total journey time. Shared EM PDR and throughput remain valid alert-delivery measures; they cannot establish differences between routing algorithms. Response labels show three decimals to reveal small numerical differences, while SUMO movement still updates every 500 ms.')

    review = ROOT / 'results/client_review/baseline_comparisons.csv'
    if review.exists():
        heading(doc, 'Direct comparison with the Fog/Cloud baseline')
        doc.add_paragraph('Fog/Cloud is the baseline and Mist Dynamic with Fog fallback is the proposed framework. The following tables show matched-seed response, signal waiting, and two supporting metrics: excess travel delay above free flow and Fog route requests. Lower is better. Missing arrivals are excluded from time metrics; request counts include every scheduled run. Excess delay is related to response time and is not independent evidence of a second journey mechanism.')
        rows = read(review)
        for context in ('Ordinary traffic', 'Controlled obstruction'):
            heading(doc, context, 2)
            for metric, label in (('ev_response_s', 'EV response time (s)'), ('traffic_light_wait_s', 'EV signal waiting (s)'), ('ev_delay_vs_freeflow_s', 'Excess travel delay (s)'), ('fog_route_requests', 'Fog route requests (count/run)')):
                doc.add_paragraph(label)
                group = [r for r in rows if r['context'] == context and r['metric'] == metric and r['config_A'] == 'MistDynamicFogFallback']
                table(doc, ('Traffic', 'Fog baseline', 'Proposed', 'Reduction %', 'Paired difference 95% interval'), [
                    (r['density'].capitalize(), f"{float(r['mean_B']):.3f}", f"{float(r['mean_A']):.3f}", f"{float(r['reduction_percent']):.2f}" if r['reduction_percent'] else 'N/A', f"{float(r['ci95_lower']):.3f} to {float(r['ci95_upper']):.3f}") for r in group])
        doc.add_paragraph('An interval wholly below zero supports a lower proposed value; an interval crossing zero is inconclusive. The selected 20 percent response target is exceeded under controlled obstruction at every density, but not in ordinary traffic. The target was selected after the measurements, so these comparisons remain exploratory. The controlled road-obstruction results do not establish the same advantage on every road. The comparison includes live routing and advance signal requests, rather than A* computation alone.')

    heading(doc, "Conclusion and limits")
    doc.add_paragraph(
        "The experiment compares alert delivery, route computation, and vehicle movement as "
        "separate parts of the response. Dynamic routing uses current road reports and only "
        "replaces a route when its stability rules are met. The controlled stall tests show "
        "how Fog recovery behaves under a defined fault. The response and waiting tables "
        "show the measured journey outcomes rather than assuming every route change saves time."
    )
    doc.add_paragraph(
        "These findings apply to the simulated 4 by 4 grid, three RSUs, 400 m radio setting, "
        "traffic scenarios, and controlled stall model. They should be tested "
        "on other road layouts and with measured onboard processing times before being "
        "generalized to real emergency deployments."
    )

    heading(doc, "How to read every graph")
    doc.add_paragraph(
        "Each bar is a mean for one configuration at one density. Error bars show the 95% "
        "confidence interval; the method used for each metric is specified in docs/METHODOLOGY.md. "
        "PDR and EM throughput include all 30 scheduled runs. Delivery-dependent measures "
        "use only runs with a delivered emergency message. A graph can contain fewer than "
        "four bars when a metric applies only to dynamic routing or fallback. Only the "
        "main waiting graphs add a fifth bar for the no-preemption control. The separate "
        "red-signal waiting graphs also include that control."
    )
    doc.add_paragraph(
        f"The following gallery contains all {len(blocks)} result graphs."
    )

    doc.add_page_break()
    heading(doc, "Complete graph gallery")
    previous = None
    for number, (key, label, explanation, path, caption) in enumerate(blocks, 1):
        if key != previous:
            heading(doc, label, 2)
            doc.add_paragraph(explanation).paragraph_format.keep_with_next = True
            previous = key
        figure = doc.add_picture(str(path), width=Inches(6.1))
        figure._inline.docPr.set("descr", f"{label}. {caption}")
        figure_paragraph = doc.paragraphs[-1]
        figure_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        figure_paragraph.paragraph_format.keep_with_next = True
        description = doc.add_paragraph(f"Figure {number}. {caption}")
        description.paragraph_format.space_after = Pt(10)
        description.paragraph_format.keep_together = True
    doc.add_paragraph(
        "Detailed metric definitions and model assumptions are in METHODOLOGY.md. "
        "The processed data and individual graph files accompany this report."
    )

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    shutil.copy2(OUTPUT, ROOT / "docs" / OUTPUT.name)
    write_graph_guide(blocks)
    print(OUTPUT)


if __name__ == "__main__":
    main()
