"""Update the submission from generated CSVs; retain the historical body."""
import csv
import json
from pathlib import Path
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from generate_matched_submission import chart

ROOT=Path(__file__).resolve().parents[1]
DATA=ROOT/'results/next_validation/completed_pilots'
FINAL=ROOT/'docs/Emergency_Vehicle_Navigation_Final_Submission.docx'
HISTORY=ROOT/'docs/Emergency_Vehicle_Navigation_Historical_e3178f2.docx'
OUT=ROOT/'scratch/completed_submission.docx'

def read(path):
    with path.open(newline='',encoding='utf-8') as stream:return list(csv.DictReader(stream))

def table(doc,headers,rows):
    t=doc.add_table(rows=1,cols=len(headers));t.style='Light Shading Accent 1'
    for cell,text in zip(t.rows[0].cells,headers):
        cell.text=text
        shading=OxmlElement('w:shd');shading.set(qn('w:fill'),'E7E6E6');cell._tc.get_or_add_tcPr().append(shading)
    repeated=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(repeated)
    for row in rows:
        for cell,value in zip(t.add_row().cells,row):cell.text=str(value)
    for row in t.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                p.paragraph_format.space_after=Pt(4)
                for run in p.runs:run.font.size=Pt(9);run.font.color.rgb=RGBColor(0,0,0)
    doc.add_paragraph()

def heading(doc,text,level=1):
    p=doc.add_heading(text,level)
    for run in p.runs:run.font.color.rgb=RGBColor(0,0,0)

def f(value):return f'{float(value):.3f}' if value else 'undefined'

def main():
    audit=json.loads((DATA/'audit.json').read_text());assert audit['passed'] and audit['complete_runs']==109
    doc=Document(HISTORY if HISTORY.exists() else FINAL)
    body=doc._element.body;original=list(body)
    title=doc.add_paragraph('Emergency Vehicle Navigation Evidence Update',style='Title')
    for run in title.runs:run.font.color.rgb=RGBColor(0,0,0)
    doc.add_paragraph('Completed pilot results after cancellation')
    doc.add_paragraph('11 October 2026')
    doc.add_paragraph('The latest experiment was cancelled before reported seeds 1 to 30 began. This update includes 109 completed held-out pilot runs only. Eleven unfinished or unstarted runs are excluded from all updated results. These pilot observations cannot establish a final performance claim.')
    doc.add_paragraph('The fair main comparison is FogCloudAStarAdvance versus MistDynamicFogFallback, with the same advance signal requests. FogCloudDynamicAdvance is a remote dynamic routing control. The preserved reported matched results remain the strongest complete evidence available.')
    doc.add_paragraph('The historical submission follows this update unchanged. Its original-policy graphs include different signal request timing, so they must not be interpreted as isolating the local computing benefit. This update and POLICY_VALIDATION.md give the matched-policy interpretation.')
    doc.add_paragraph('No simulations were restarted. All CSVs and new figures were regenerated from fully completed raw evidence. The raw files from unfinished runs are preserved separately and are not included.')
    doc.add_page_break()
    for context in ('ordinary','obstruction'):
        heading(doc,'Preserved matched '+context+' evidence')
        pairs=[r for r in read(ROOT/f'results/policy_validation/{context}/paired_comparisons.csv')
               if r['comparison']=='matched_policy' and r['metric'] in ('route_decision_ms','ev_response_s','traffic_light_wait_s')]
        labels={'route_decision_ms':'A Decision ms','ev_response_s':'B Response s','traffic_light_wait_s':'D Signal wait s'}
        table(doc,['Density and metric','Baseline','Proposed','Change %','Paired 95% interval','Pairs'],
              [[r['density']+' '+labels[r['metric']],f(r['mean_B']),f(r['mean_A']),f(r['percent_change']),
                '['+f(r['ci95_lower'])+', '+f(r['ci95_upper'])+']',r['n_pairs']] for r in pairs])
        doc.add_paragraph('Intervals are paired proposed minus baseline 95 percent intervals, retaining the existing analysis method. Negative time differences favor proposed. Zero baseline waiting gives an undefined percent change. The inference methods remain documented in METHODOLOGY.md.')
        doc.add_paragraph('Ordinary response improves by only 0.3 to 0.7 percent and does not meet the 20 percent reference. Fixed obstruction response improves by roughly 40 to 45 percent. Extra supported fixed-obstruction improvements are full stops and matched background time loss. Signal waiting remains unestablished: ordinary waiting is zero for both, and fixed-obstruction difference intervals include zero. A lower nonzero waiting requirement is therefore unmet.')
        doc.add_page_break()
    heading(doc,'Completed pilot coverage')
    coverage=read(DATA/'coverage.csv')
    names={'FogCloudAStarAdvance':'Baseline','MistDynamicFogFallback':'Proposed','FogCloudDynamicAdvance':'Dynamic Fog'}
    table(doc,['Window s','Trips','Framework','Complete','Scheduled'],
          [[r['pilot_window_s'],r['density'].replace('heavy',''),names[r['configuration']],r['completed_runs'],10] for r in coverage])
    doc.add_paragraph('Each pilot cell scheduled all seeds 101 to 110. The first pilot completed 60 of 60 runs. The second completed 49 of 60. The coverage CSV lists exactly which seeds are included and excluded. Comparisons use only matched completed pairs; completion patterns can bias the incomplete second pilot.')
    doc.add_paragraph('The first pilot formed queues but baseline response rose only from 148.25 s to 150.00 s between 400 and 600 trips, a paired increase of 1.75 s. This missed the predeclared 5 s criterion. The second pilot concentrated departures into 120 seconds, but was cancelled before completion; it has no final congestion qualification.')
    doc.add_page_break()
    labels={'route_decision_ms':'A Decision ms','ev_response_s':'B Response s','ev_full_stops':'C Full stops',
            'ev_green_approach_fraction':'C Green share','background_time_loss_s':'C Background loss s','traffic_light_wait_s':'D Signal wait s'}
    pairs=read(DATA/'paired_comparisons.csv')
    for window in (360,120):
        heading(doc,f'Completed pilot measurements with {window} second departures')
        chosen=[r for r in pairs if r['contrast']=='main_matched' and int(r['pilot_window_s'])==window and r['metric'] in labels]
        table(doc,['Trips and metric','Baseline','Proposed','Change %','Paired 95% interval','Pairs'],
          [[r['density'].replace('heavy','')+' '+labels[r['metric']],f(r['baseline_mean']),f(r['proposed_mean']),f(r['percent_change']),
            '['+f(r['ci95_lower'])+', '+f(r['ci95_upper'])+']',r['n_pairs']] for r in chosen])
        doc.add_paragraph('These are descriptive held-out pilot comparisons, not final reported results. Intervals use paired bootstrap resampling with 10000 draws and fixed seed 42. Every included run completed the full 900 second horizon. Unfinished runs contribute no values.')
        doc.add_page_break()
    heading(doc,'Measurement and interpretation')
    doc.add_paragraph('SUMO generates shared traffic and moves the vehicles. Veins and OMNeT++ carry emergency messages and Fog requests. The matched frameworks use identical 250 m or 20 s signal requests and safety rules. The watchdog stays at 500 ms, telemetry validation at 0 ms and polling at 100 ms, with underlying SUMO states every 500 ms.')
    doc.add_paragraph('The 900 ms Mist stalls are controlled injected faults. Existing Mist periodic reviews compute synchronously, while Fog reviews include modeled remote service and radio delay. This ablation is not a physical OBU processor benchmark. The cancelled pilot has no completed 600-trip dynamic Fog runs; no value or comparison is invented for that missing cell.')
    doc.add_paragraph('Signal waiting measures speed below 0.1 m/s within 20 m of the next signal. Whole-route time below 1 m/s detects slow crawling elsewhere. Whole-route stops also include the shared scheduled start stop. Zero waiting remains zero when measured. Mean speed depends on journey time. Background loss uses common observed vehicle IDs and is not a pure causal measure of preemption harm.')
    doc.add_paragraph('The preferred seven metrics and all available secondary metrics are in the generated CSV summaries. All completed-pilot graph panels are also in COMPLETED_PILOT_RESULTS.md. The following graphs show the main comparison and available dynamic Fog ablation, with 95 percent intervals. Missing cells are left missing.')
    folder=ROOT/'scratch/completed_word_graphs';folder.mkdir(exist_ok=True)
    metrics={'route_decision_ms':('Route decision latency','ms'),'ev_response_s':('EV response time','s'),
        'traffic_light_wait_s':('EV signal waiting','s'),'ev_full_stops':('EV full stops','stops/run'),
        'background_time_loss_s':('Background time loss','s/vehicle'),'ev_green_approach_fraction':('Green approach share','share')}
    for window in (360,120):
        summary=read(DATA/f'pilot_{window}/summary.csv')
        for index,(metric,(label,unit)) in enumerate(metrics.items()):
            if index%2==0:doc.add_page_break();heading(doc,f'Completed {window} second pilot graphs')
            image=chart(summary,metric,label,unit,folder,f'Pilot {window} s')
            doc.add_picture(str(image),width=Inches(6.5))
            doc.add_paragraph(label+' from completed runs only. Error bars show 95 percent intervals.')
    doc.add_page_break();heading(doc,'Historical submission follows')
    doc.add_paragraph('The following preserved material documents earlier complete experiments. Original-policy results do not hold signal timing equal. Use the matched-policy tables at the start of this update for the fair headline comparison. The cancelled pilot does not replace or complete those experiments.')
    doc.add_page_break()
    added=[e for e in list(body) if e not in original]
    for element in reversed(added):body.insert(0,element)
    OUT.parent.mkdir(exist_ok=True);doc.save(OUT)
    print(OUT)

if __name__=='__main__':main()
