"""Check the integrated document against its published evidence sources."""
import csv
import unittest
from zipfile import ZipFile
from lxml import etree
from pathlib import Path
from docx import Document

ROOT=Path(__file__).resolve().parents[1]

class TestCompleteProjectReport(unittest.TestCase):
    def test_full_report_and_source_values(self):
        path=ROOT/'docs/Emergency_Vehicle_Navigation_Final_Submission.docx'
        doc=Document(path)
        text='\n'.join(p.text for p in doc.paragraphs)
        for heading in ('Project overview','System design','How an emergency is handled',
                        'Routing frameworks','Experimental design','Seven primary metrics',
                        'Conclusions and remaining requirements'):
            self.assertIn(heading,text)
        self.assertNotIn('Historical submission follows',text)
        self.assertEqual(len(doc.inline_shapes),100)
        cells='\n'.join(c.text for t in doc.tables for row in t.rows for c in row.cells)
        for context in ('ordinary','obstruction'):
            with (ROOT/f'results/policy_validation/{context}/paired_comparisons.csv').open(newline='') as stream:
                rows=list(csv.DictReader(stream))
            for row in rows:
                if row['comparison']=='matched_policy' and row['metric'] in ('ev_response_s','traffic_light_wait_s'):
                    for key in ('mean_A','mean_B','ci95_lower','ci95_upper'):
                        self.assertIn(f"{float(row[key]):.3f}",cells)
        self.assertIn('109 fully completed',text)
        self.assertRegex(text,r'11 (?:unfinished|cases that were not completed or started)')
        # Every graph must resolve to its own embedded asset after Word edits.
        with ZipFile(path) as package:
            ns={'a':'http://schemas.openxmlformats.org/drawingml/2006/main',
                'r':'http://schemas.openxmlformats.org/officeDocument/2006/relationships'}
            root=etree.fromstring(package.read('word/document.xml'))
            rels={r.get('Id'):r.get('Target') for r in etree.fromstring(
                package.read('word/_rels/document.xml.rels'))}
            images=[package.read('word/'+rels[b.get('{'+ns['r']+'}embed')])
                    for b in root.findall('.//a:blip',ns)]
            self.assertEqual(len(images),100)
            self.assertEqual(len(set(images)),100)

if __name__=='__main__':unittest.main()
